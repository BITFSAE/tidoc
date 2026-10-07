"""Actual docxtpl, sandbox, all Word parts and source/executable IPC regression."""
from copy import deepcopy
from pathlib import Path
import json
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tidoc.services.export_context import build_export_context
from tidoc_print.template_validation import validate_template
from tidoc_print.template_renderer import render_template, upper_rmb, TemplateError
from tidoc_print.protocol import validate_request, render_request, ProtocolError, capabilities
from tidoc_print.context import component_resource_path, context_field_catalog


def context():
    return build_export_context([{'id':'one','title':'学校','title_profile_id':'school','total':'0','fields':{'paid_amount':{'current':'0'}},'items':[{'name':'长中文<&>商品','unit':'个','quantity':'1','total':'0'}]}],definition={'scheme':{'organization':{'name':'实验室','department':'研发','purpose':'设备采购'}}},options={'date':'2026-10-06'})


def make_template(path,body='{{ title.name }}',header=None,footer=None):
    document=Document()
    document.add_paragraph(body)
    if header:document.sections[0].header.paragraphs[0].text=header
    if footer:document.sections[0].footer.paragraphs[0].text=footer
    document.save(path)
    return path


def test_whole_xml_structural_loops_headers_footers_escape_and_static_image(tmp_path):
    from PIL import Image
    document=Document()
    document.add_paragraph('{%p if title %}')
    document.add_paragraph('{{ title.name }}')
    document.add_paragraph('{%p endif %}')
    table=document.add_table(rows=4,cols=2)
    table.cell(0,0).text='名称';table.cell(0,1).text='金额'
    table.cell(1,0).text='{%tr for row in rows %}'
    table.cell(2,0).text='{{ loop.index }} {{ row.product_name }}'
    table.cell(2,1).text='{{ row.total|money }}'
    table.cell(3,0).text='{%tr endfor %}'
    document.sections[0].header.paragraphs[0].text='{{ scheme.organization.name }}'
    document.sections[0].footer.paragraphs[0].text='{{ export.date|date("chinese") }}'
    image=tmp_path/'logo.png';Image.new('RGB',(20,20),'blue').save(image)
    document.add_picture(str(image))
    template=tmp_path/'template.docx';document.save(template)
    assert validate_template(template,context=context())==[]
    out=render_template(template,context(),tmp_path/'output.docx')
    doc=Document(out)
    assert len(doc.tables[0].rows)==2
    assert doc.tables[0].cell(1,0).text=='1 长中文<&>商品'
    assert doc.sections[0].header.paragraphs[0].text=='实验室'
    assert doc.sections[0].footer.paragraphs[0].text=='2026年10月6日'
    assert len(doc.inline_shapes)==1


@pytest.mark.parametrize('syntax',[
    '{{ title.__class__ }}','{{ title.name.upper() }}','{{ range(3) }}',
    '{% set x = 1 %}','{% include "other" %}','{% macro x() %}x{% endmacro %}',
    '{{ totals.invoice + 1 }}','{{ title.name|safe }}','{{ title[export.output_id] }}',
    '{% for key in title %}{{ key }}{% endfor %}',
    '{% for a in rows %}{% for b in rows %}{% for c in rows %}{% for d in rows %}x{% endfor %}{% endfor %}{% endfor %}{% endfor %}',
])
def test_unsafe_ast_rejected(tmp_path,syntax):
    template=make_template(tmp_path/'bad.docx',syntax)
    diagnostics=validate_template(template,context=context())
    assert diagnostics
    assert diagnostics[0]['file'].endswith('bad.docx')
    with pytest.raises(TemplateError):render_template(template,context(),tmp_path/'out.docx')
    assert not (tmp_path/'out.docx').exists()


def test_undefined_even_with_default_and_in_footer(tmp_path):
    template=make_template(tmp_path/'bad.docx','{{ title.name }}',footer='{{ payee.bank_number|default("") }}')
    diagnostics=validate_template(template,context=context())
    assert any(diag['code']=='UNKNOWN_CONTEXT_FIELD' and 'footer' in diag['location'] for diag in diagnostics)
    with pytest.raises(TemplateError):render_template(template,context(),tmp_path/'out.docx')


def test_split_runs_rejected(tmp_path):
    document=Document();p=document.add_paragraph();p.add_run('{{ title.');p.add_run('name }}').bold=True
    template=tmp_path/'split.docx';document.save(template)
    assert any(d['code']=='SPLIT_TEMPLATE_TAG' for d in validate_template(template))


def test_external_resources_fields_ole_and_unsupported_parts(tmp_path):
    document=Document();document.add_paragraph('x')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'INCLUDETEXT "file:///secret"')
    document.paragraphs[0]._p.append(field)
    path=tmp_path/'field.docx';document.save(path)
    assert any(d['code']=='AUTOMATIC_FIELD' for d in validate_template(path))
    original=make_template(tmp_path/'normal.docx')
    bad=tmp_path/'external.docx'
    with zipfile.ZipFile(original) as src,zipfile.ZipFile(bad,'w') as dest:
        for name in src.namelist():
            content=src.read(name)
            if name=='word/_rels/document.xml.rels':
                content=content.replace(b'</Relationships>',b'<Relationship Id="remote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://example.com/image" TargetMode="External"/></Relationships>')
            dest.writestr(name,content)
        dest.writestr('word/embeddings/object.bin',b'x')
    codes={diag['code'] for diag in validate_template(bad)}
    assert {'EXTERNAL_DOCX_RESOURCE','UNSAFE_DOCX_PART'}<=codes


def test_money_rmb_missing_negative_and_front_zero():
    assert upper_rmb('0')=='零元整'
    assert upper_rmb('-1001.05')=='负壹仟零壹元零伍分'
    assert upper_rmb(None)==''
    assert upper_rmb('100010001')=='壹亿零壹万零壹元整'


def request(tmp_path):
    root=tmp_path/'resources';root.mkdir()
    template=make_template(root/'doc.docx','{{ rows["bad"] }}')
    make_template(template,'{{ title.name }} {{ totals.invoice|money }}')
    return {'ipc_version':2,'job_id':'j','resources_root':str(root),'output_dir':str(tmp_path/'out'),'files':[{'output':{'id':'doc','type':'docx','payee_mode':'none'},'context':context(),'template':'doc.docx','filename':'export.docx'}]}


def test_source_and_executable_same_v2_protocol_and_semantics(tmp_path):
    req=request(tmp_path)
    direct=render_request(req)
    assert Path(direct['files'][0]['path']).exists()
    req['output_dir']=str(tmp_path/'external')
    input_path=tmp_path/'input.json';output_path=tmp_path/'result.json'
    input_path.write_text(json.dumps(req,ensure_ascii=False),'utf-8')
    proc=subprocess.run([sys.executable,'-m','tidoc_print','--input',str(input_path),'--result',str(output_path)],text=True,encoding='utf-8',capture_output=True,timeout=20)
    assert proc.returncode==0,proc.stderr
    result=json.loads(output_path.read_text('utf-8'))
    assert result['ok']
    assert Document(result['data']['files'][0]['path']).paragraphs[0].text==Document(direct['files'][0]['path']).paragraphs[0].text
    cap=subprocess.run([sys.executable,'-m','tidoc_print','--capabilities'],text=True,encoding='utf-8',capture_output=True,timeout=10)
    assert json.loads(cap.stdout)['ipc_versions']==[1,2]


def test_v2_isolation_paths_pure_data_and_cancellation(tmp_path):
    req=request(tmp_path)
    unsafe=deepcopy(req);unsafe['files'][0]['template']='../escape.docx'
    with pytest.raises(ProtocolError):validate_request(unsafe)
    unsafe=deepcopy(req);unsafe['files'][0]['filename']='../escape.docx'
    with pytest.raises(ProtocolError):validate_request(unsafe)
    unsafe=deepcopy(req);unsafe['files'][0]['context']['entries'].append({**unsafe['files'][0]['context']['entries'][0],'title':{**unsafe['files'][0]['context']['entries'][0]['title'],'name':'other','tax_id':'other'}})
    with pytest.raises(ProtocolError,match='抬头隔离'):validate_request(unsafe)
    with pytest.raises(ProtocolError,match='取消'):render_request(req,cancel_check=lambda:True)
    assert not Path(req['output_dir']).exists()


def test_legacy_ipc_converts_to_same_docxtpl_renderer(tmp_path):
    request={'entries':[{'entry_id':'a','title':'学校','invoice_no':'1','total':'12.30','items':[{'actual_name':'物品','product_name':'品名','quantity':'1','unit':'个','total':'12.30'}]}],'out_dir':str(tmp_path/'legacy'),'options':{'make_entry_bundle_pdf':False,'make_acceptance_doc':False},'profiles':{}}
    normalized=validate_request(request)
    assert normalized['ipc_version']==2
    result=render_request(normalized)
    doc=Document(result['files'][0]['path'])
    assert doc.tables[0].cell(1,0).text=='物品'
    assert doc.tables[0].cell(1,1).text=='品名'
    assert doc.tables[0].cell(1,2).text=='¥12.30'


def test_public_catalog_static_and_runtime_loop_fields_agree(tmp_path):
    template=make_template(tmp_path/'public.docx','{% for entry in entries %}{{ entry.actual_name }} {{ entry.product_name }} {{ entry.invoice_no }}{% endfor %}{% for row in rows %}{{ row.product_name }} {{ row.is_first }}{% endfor %}')
    assert validate_template(template)==[]
    assert validate_template(template,context=context())==[]
    out=render_template(template,context(),tmp_path/'public-out.docx')
    assert '长中文<&>商品' in Document(out).paragraphs[0].text


def test_material_pdf_all_custom_roles_layout_a4_and_original_unchanged(tmp_path):
    from PIL import Image
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas
    from tidoc_print.pdf_merge import render_pdf_bundle
    root=tmp_path/'materials';root.mkdir()
    invoice=root/'invoice.pdf';c=canvas.Canvas(str(invoice),pagesize=(300,400));c.drawString(20,300,'INVOICE');c.save()
    resources=[{'id':'invoice','entry_id':'one','role_id':'invoice','path':'invoice.pdf'}]
    for i in range(5):
        path=root/f'image-{i}.png';Image.new('RGB',(80,160),'red').save(path)
        resources.append({'id':str(i),'entry_id':'one','role_id':'physical_image','path':path.name})
    custom=root/'custom.pdf';c=canvas.Canvas(str(custom));c.drawString(20,500,'CUSTOM');c.save()
    resources.append({'id':'contract','entry_id':'one','role_id':'custom:test:contract','path':'custom.pdf'})
    out=render_pdf_bundle(context(),{'pdf':{'include_roles':['invoice','physical_image','custom:test:contract'],'content_order':'role','image_layout':'a4_portrait_4','page_size':'a4','numbering':True}},resources,root,tmp_path/'bundle.pdf')
    pages=PdfReader(out).pages
    assert len(pages)==4
    assert all(abs(float(page.mediabox.width)-595.275)<1 for page in pages)
    assert 'INVOICE' in pages[0].extract_text()
    assert 'CUSTOM' in pages[-1].extract_text()
    assert 'No.' not in PdfReader(invoice).pages[0].extract_text()
    assert float(PdfReader(invoice).pages[0].mediabox.width)==300
    (root/'unsupported.txt').write_text('raw','utf-8')
    with pytest.raises(ValueError,match='不能转换'):
        render_pdf_bundle(context(),{'pdf':{'include_roles':['custom:test:contract']}},[{'entry_id':'one','role_id':'custom:test:contract','path':'unsupported.txt'}],root,tmp_path/'invalid.pdf')


def test_adapter_and_print_catalogs_share_types_and_extension_aliases():
    from tidoc.adapters.registry import FIELD_CATALOG, FILENAME_FIELDS, field_catalog
    from tidoc_print.context import filename_field_catalog, format_filename
    definition={'fields':[{'scope':scope,'id':'code','type':'text'}
                          for scope in ('scheme','payee','entry','batch','export')]}
    assert FIELD_CATALOG == context_field_catalog()
    assert field_catalog(definition) == context_field_catalog(definition)
    assert tuple(FILENAME_FIELDS) == filename_field_catalog()
    assert format_filename('{entry.invoice_no}_{role.label}',{'entry':{'invoice_no':'001'},'role':{'label':'付款'}})=='001_付款'
    for pattern in ('{payee.account_number}', '{material.name}', '{entry.total}'):
        with pytest.raises(ValueError):format_filename(pattern,{})
    catalog=context_field_catalog(definition)
    assert catalog['row.claimant'] == catalog['rows[].claimant'] == 'text'
    assert catalog['row.entry.claimant'] == 'object'
    assert catalog['entry.fields.code'] == catalog['entries[].fields.code'] == 'text'
    assert catalog['row.entry.batch.fields.code'] == 'text'
    catalog['title.name']='integer'
    assert context_field_catalog()['title.name']=='text'


def test_catalog_and_resource_lookup_import_no_core_or_renderer_modules():
    code='''
import builtins
original=builtins.__import__
def isolated(name,*args,**kwargs):
    if name.startswith(('tidoc.db','tidoc.api','webview','docx','docxtpl','jinja2','jsonschema')):
        raise AssertionError('heavy import: '+name)
    return original(name,*args,**kwargs)
builtins.__import__=isolated
from tidoc_print.context import context_field_catalog
from tidoc.adapters.registry import FIELD_CATALOG
assert context_field_catalog()==FIELD_CATALOG
'''
    result=subprocess.run([sys.executable,'-c',code],capture_output=True,text=True,encoding='utf-8',timeout=10)
    assert result.returncode==0,result.stderr


def frozen_resources(tmp_path,monkeypatch):
    import shutil
    root=tmp_path/'frozen'
    resources=[('tidoc_print','context_fields.json'),
               ('schemas','team-adapter/1/context.schema.json'),
               ('tidoc_print','templates/报账说明模板.docx'),
               ('tidoc_print','templates/验收单模板.docx'),
               ('tidoc','builtin_adapters/org.bitfsae.reimbursement/scheme.json')]
    for package,relative in resources:
        source=component_resource_path(relative,package=package)
        target=root/package/relative;target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copyfile(source,target)
    monkeypatch.setattr(sys,'frozen',True,raising=False)
    monkeypatch.setattr(sys,'_MEIPASS',str(root),raising=False)
    return root


def legacy_request(tmp_path):
    return {'ipc_version':1,'job_id':'legacy-fixture','entries':[{
        'entry_id':'a','title':'学校','total':'-12.30','paid_amount':'0',
        'items':[{'actual_name':'物品','product_name':'品名','quantity':'1','unit':'个','total':'-12.30'}]
    }],'out_dir':str(tmp_path/'legacy-out'),
        'options':{'make_entry_bundle_pdf':False,'document_date':'2026-10-06'},'profiles':{}}


def test_frozen_catalog_and_v1_defaults_use_packaged_resources(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    assert component_resource_path('context_fields.json')==root/'tidoc_print/context_fields.json'
    normalized=validate_request(legacy_request(tmp_path))
    assert all((Path(normalized['resources_root'])/item['template']).is_relative_to(root) for item in normalized['files'])
    assert normalized['files'][0]['context']['scheme']['organization']['department']=='机械与车辆学院'
    result=render_request(normalized)
    assert len(result['files'])==2
    assert capabilities()['ipc_versions']==[1,2]


def test_missing_frozen_catalog_is_reported_without_source_fallback(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    (root/'tidoc_print/context_fields.json').unlink()
    with pytest.raises(FileNotFoundError,match='context_fields.json'):context_field_catalog()
    status=capabilities()
    assert 'resource:tidoc_print/context_fields.json' in status['missing']
    assert 'docx' not in status['renderers']
    assert 'pdf_bundle' in status['renderers']
    template=make_template(tmp_path/'template.docx')
    assert validate_template(template)[0]['code']=='MISSING_COMPONENT_RESOURCE'


def test_missing_frozen_legacy_scheme_is_explicit_and_v2_still_works(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    (root/'tidoc/builtin_adapters/org.bitfsae.reimbursement/scheme.json').unlink()
    with pytest.raises(ProtocolError) as error:validate_request(legacy_request(tmp_path))
    assert error.value.code=='MISSING_COMPONENT_RESOURCE'
    assert capabilities()['ipc_versions']==[2]
    assert 'docx' in capabilities()['renderers']
    assert Path(render_request(request(tmp_path))['files'][0]['path']).exists()


def test_invalid_packaged_catalog_is_not_misreported_as_bad_template(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    (root/'tidoc_print/context_fields.json').write_text('{"title.name":"unknown"}','utf-8')
    template=make_template(tmp_path/'template.docx')
    assert validate_template(template)[0]['code']=='INVALID_CONTEXT_CATALOG'


def test_same_catalog_checks_properties_headers_footers_and_footnotes(tmp_path):
    document=Document();document.add_paragraph('{{ title.name }}')
    document.core_properties.title='{{ title.name }}'
    document.sections[0].header.paragraphs[0].text='{{ totals.unknown }}'
    document.sections[0].footer.paragraphs[0].text='{{ totals.unknown }}'
    source=tmp_path/'source.docx';document.save(source)
    template=tmp_path/'parts.docx'
    with zipfile.ZipFile(source) as src,zipfile.ZipFile(template,'w') as dest:
        for name in src.namelist():dest.writestr(name,src.read(name))
        dest.writestr('word/footnotes.xml','''<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>{{ totals.unknown }}</w:t></w:r></w:p></w:footnote></w:footnotes>''')
    diagnostics=validate_template(template)
    locations=[d['location'] for d in diagnostics if d['code']=='UNKNOWN_CONTEXT_FIELD']
    assert any('header' in loc for loc in locations)
    assert any('footer' in loc for loc in locations)
    assert any('footnotes' in loc for loc in locations)
    assert all('core.xml' not in loc for loc in locations)


def test_declared_extension_loops_and_scalar_boundaries(tmp_path):
    definition={'fields':[{'scope':scope,'id':'code','type':'text'} for scope in ('entry','batch','export')]}
    template=make_template(tmp_path/'fields.docx','{% for entry in entries %}{{ entry.fields.code }} {{ entry.batch.fields.code }}{% endfor %}{% for row in rows %}{{ row.entry.fields.code }}{% endfor %}{{ export.fields.code }}')
    assert validate_template(template,definition=definition)==[]
    assert any(d['code']=='UNKNOWN_CONTEXT_FIELD' for d in validate_template(template))
    invalid=make_template(tmp_path/'scalar.docx','{% for row in rows %}{{ row.claimant.name }}{% endfor %}')
    assert any(d['code']=='UNKNOWN_CONTEXT_FIELD' for d in validate_template(invalid))


def test_arbitrary_three_word_output_ids_source_and_cli_match(tmp_path):
    req=request(tmp_path);base=req['files'][0]
    req['files']=[]
    for oid,label in [('purchase_request','采购申请'),('settlement','结算说明'),('equipment_record','设备登记')]:
        item=deepcopy(base);item['output'].update(id=oid,label=label)
        item['context']['export']['output_id']=oid
        item['template']=oid+'.docx';item['filename']=label+'.docx'
        make_template(Path(req['resources_root'])/item['template'],label+' {{ title.name }} {{ export.output_id }}')
        req['files'].append(item)
    direct=render_request(req)
    req['output_dir']=str(tmp_path/'cli-out')
    result_path=run_cli(tmp_path,req)
    response=json.loads(result_path.read_text('utf-8'))
    assert response['ok'],response
    assert len(direct['files'])==len(response['data']['files'])==3
    assert [f['output_id'] for f in response['data']['files']]==['purchase_request','settlement','equipment_record']
    for expected,actual in zip(direct['files'],response['data']['files']):
        assert Document(expected['path']).paragraphs[0].text==Document(actual['path']).paragraphs[0].text


def run_cli(tmp_path,payload,*args):
    input_path=tmp_path/'cli-request.json';result_path=tmp_path/'cli-result.json'
    input_path.write_text(json.dumps(payload,ensure_ascii=False),'utf-8')
    proc=subprocess.run([sys.executable,'-m','tidoc_print','--input',str(input_path),'--result',str(result_path),*args],capture_output=True,text=True,encoding='utf-8',timeout=20)
    assert result_path.exists(),proc.stderr
    assert (proc.returncode==0)==json.loads(result_path.read_text('utf-8'))['ok']
    return result_path


def test_v1_cli_and_source_share_defaults_and_semantics(tmp_path):
    payload=legacy_request(tmp_path)
    direct=render_request(payload)
    payload['out_dir']=str(tmp_path/'cli-out')
    response=json.loads(run_cli(tmp_path,payload).read_text('utf-8'))
    assert response['ok'],response
    assert [f['output_id'] for f in direct['files']]==['reimburse_doc','acceptance_doc']
    for expected,actual in zip(direct['files'],response['data']['files']):
        a,b=Document(expected['path']),Document(actual['path'])
        assert [p.text for p in a.paragraphs]==[p.text for p in b.paragraphs]
        assert [[cell.text for cell in row.cells] for row in a.tables[0].rows]==[[cell.text for cell in row.cells] for row in b.tables[0].rows]
    doc=Document(direct['files'][0]['path'])
    assert '机械与车辆学院' in ''.join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1,2).text=='¥-12.30'
    assert '工训楼' in ''.join(cell.text for row in Document(direct['files'][1]['path']).tables[0].rows for cell in row.cells)


@pytest.mark.parametrize('version',[1,2])
@pytest.mark.parametrize('failure',['cancel','timeout'])
def test_cli_cancel_timeout_for_both_protocols_never_publish(tmp_path,version,failure):
    req=legacy_request(tmp_path) if version==1 else request(tmp_path)
    args=[]
    if failure=='cancel':
        cancel=tmp_path/'cancel';cancel.touch();args=['--cancel-file',str(cancel)]
    else:
        args=['--timeout','0.000000001']
    response=json.loads(run_cli(tmp_path,req,*args).read_text('utf-8'))
    assert not response['ok']
    assert response['code']==('EXPORT_CANCELED' if failure=='cancel' else 'EXPORT_TIMEOUT')
    assert not Path(req.get('out_dir') or req['output_dir']).exists()
    assert not list(tmp_path.glob('.tidoc-print-*'))


@pytest.mark.parametrize('failure',['cancel','timeout'])
def test_late_cancel_timeout_removes_successfully_rendered_staging(tmp_path,monkeypatch,failure):
    import tidoc_print.template_renderer as renderer
    import tidoc_print.protocol as protocol
    req=request(tmp_path);rendered=[];clock=[0.0]
    original=renderer.render_template
    def complete_then_interrupt(*args,**kwargs):
        path=original(*args,**kwargs);rendered.append(path)
        clock[0]=121.0
        return path
    monkeypatch.setattr(renderer,'render_template',complete_then_interrupt)
    monkeypatch.setattr(protocol.time,'monotonic',lambda:clock[0])
    with pytest.raises(ProtocolError) as error:
        render_request(req,cancel_check=(lambda:bool(rendered)) if failure=='cancel' else None)
    assert error.value.code==('EXPORT_CANCELED' if failure=='cancel' else 'EXPORT_TIMEOUT')
    assert rendered and not rendered[0].exists()
    assert not Path(req['output_dir']).exists()
    assert not list(tmp_path.glob('.tidoc-print-*'))


def test_missing_and_changed_request_resources_fail_before_render(tmp_path):
    req=request(tmp_path)
    missing=deepcopy(req);missing['files'][0]['template']='missing.docx'
    with pytest.raises(ProtocolError) as error:render_request(missing)
    assert error.value.code=='MISSING_RESOURCE'
    changed=deepcopy(req);changed['files'][0]['template_sha256']='0'*64
    with pytest.raises(ProtocolError) as error:render_request(changed)
    assert error.value.code=='RESOURCE_CHANGED'
    missing=deepcopy(req);missing['files'][0]['resources']=[{'path':'missing.pdf'}]
    with pytest.raises(ProtocolError) as error:render_request(missing)
    assert error.value.code=='MISSING_RESOURCE'
    assert not Path(req['output_dir']).exists()


def test_missing_frozen_legacy_template_does_not_block_custom_v2(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    (root/'tidoc_print/templates/验收单模板.docx').unlink()
    with pytest.raises(ProtocolError) as error:validate_request(legacy_request(tmp_path))
    assert error.value.code=='MISSING_COMPONENT_RESOURCE'
    status=capabilities()
    assert status['ipc_versions']==[2]
    assert status['renderers']==['docx','pdf_bundle']
    assert 'resource:tidoc_print/templates/验收单模板.docx' in status['missing']
    assert Path(render_request(request(tmp_path))['files'][0]['path']).is_file()


@pytest.mark.parametrize('mutation,location', [
    ('missing_totals', '/files/0/context'),
    ('unknown_field', '/files/0/context'),
    ('numeric_money', '/files/0/context/entries/0/invoice/total'),
    ('bad_rows', '/files/0/context/rows'),
    ('unknown_option', '/files/0/context/export/options'),
])
def test_runtime_rejects_malformed_context_before_any_output(tmp_path,mutation,location):
    req=request(tmp_path)
    ctx=req['files'][0]['context']
    if mutation=='missing_totals':
        ctx.pop('totals')
    elif mutation=='unknown_field':
        ctx['secret_account']='PRIVATE-ACCOUNT-0000123'
    elif mutation=='numeric_money':
        ctx['entries'][0]['invoice']['total']=0
    elif mutation=='bad_rows':
        ctx['rows']='PRIVATE-ROWS'
    else:
        ctx['export']['options']['secret_account']='PRIVATE-ACCOUNT-0000123'
    with pytest.raises(ProtocolError) as error:
        render_request(req)
    assert error.value.code=='INVALID_CONTEXT'
    assert any(item['location']==location for item in error.value.diagnostics)
    assert 'PRIVATE' not in json.dumps(error.value.diagnostics)
    assert not Path(req['output_dir']).exists()
    assert not list(tmp_path.glob('.tidoc-print-*'))
    source=tmp_path/'request.json';source.write_text(json.dumps(req),'utf-8')
    result=tmp_path/'result.json'
    proc=subprocess.run([sys.executable,'-m','tidoc_print','--input',str(source),
                         '--result',str(result)],capture_output=True,text=True,encoding='utf-8',timeout=20)
    assert proc.returncode==1
    external=json.loads(result.read_text('utf-8'))
    assert external['code']==error.value.code
    assert external['diagnostics']==error.value.diagnostics
    assert not Path(req['output_dir']).exists()


def test_missing_frozen_context_schema_disables_renderers_and_self_test(tmp_path,monkeypatch):
    root=frozen_resources(tmp_path,monkeypatch)
    (root/'schemas/team-adapter/1/context.schema.json').unlink()
    status=capabilities()
    assert status['renderers']==[]
    assert status['ipc_versions']==[]
    assert status['context_versions']==[]
    assert 'resource:schemas/team-adapter/1/context.schema.json' in status['missing']
    with pytest.raises(ProtocolError) as error:
        render_request(request(tmp_path))
    assert error.value.code=='INVALID_COMPONENT_RESOURCE'
    assert not (tmp_path/'out').exists()


def test_core_component_status_retains_source_resource_failure(tmp_path,monkeypatch):
    from tidoc.services import printing
    import tidoc_print.protocol as protocol
    monkeypatch.delattr(sys,'frozen',raising=False)
    monkeypatch.setattr(protocol,'capabilities',lambda:{'renderers':[],
        'missing':['resource:schemas/team-adapter/1/context.schema.json']})
    status=printing.component_status(tmp_path)
    assert status['available'] is False
    assert status['missing']==['resource:schemas/team-adapter/1/context.schema.json']
    assert '修复' in status['error']


def test_component_self_test_detects_unloadable_dependency_despite_find_spec(monkeypatch):
    import builtins
    from tidoc_print.__main__ import self_test_status
    original=builtins.__import__
    def broken_pillow(name,*args,**kwargs):
        if name=='PIL':
            raise ImportError('simulated unloadable image library')
        return original(name,*args,**kwargs)
    monkeypatch.setattr(builtins,'__import__',broken_pillow)
    result=self_test_status()
    assert result['missing']==[]
    assert result['ok'] is False
    assert result['code']=='COMPONENT_SELF_TEST_FAILED'


def test_context_projection_excludes_batch_settings_and_local_versions():
    ctx=build_export_context([{'id':'one','title':'学校','total':'0',
        'batch':{'id':'batch','name':'批次','note':None,'updated_at':'local-version',
                 'output_settings':{'print.numbering':False}}}])
    from tidoc_print.context_validation import validate_context
    assert validate_context(ctx)==[]
    assert ctx['batch']==ctx['entries'][0]['batch']==ctx['rows'][0]['entry']['batch']
    assert set(ctx['batch'])=={'id','name','notes','fields'}
    assert ctx['batch']['notes']==''


@pytest.mark.parametrize('timeout',[False,None,'1',0,-1,121,float('nan'),float('inf')])
def test_invalid_timeout_never_creates_output(tmp_path,timeout):
    req=request(tmp_path);req['timeout_seconds']=timeout
    with pytest.raises(ProtocolError):render_request(req)
    assert not Path(req['output_dir']).exists()


def test_runtime_context_cannot_expand_public_field_permissions(tmp_path):
    ctx=context();ctx['title']['secret']='private';ctx['export']['options']['payee_ids_by_binding']={'other':'private'}
    template=make_template(tmp_path/'undeclared.docx','{{ title.secret }} {{ export.options.payee_ids_by_binding }}')
    assert any(d['code']=='UNKNOWN_CONTEXT_FIELD' for d in validate_template(template,context=ctx))
    with pytest.raises(TemplateError):render_template(template,ctx,tmp_path/'never.docx')


def test_declared_but_missing_runtime_field_still_uses_strict_undefined(tmp_path):
    from jinja2 import UndefinedError
    definition={'fields':[{'scope':'export','id':'code','type':'text'}]}
    template=make_template(tmp_path/'declared.docx','{{ export.fields.code|default("fallback") }}')
    assert validate_template(template,definition=definition)==[]
    with pytest.raises(UndefinedError):render_template(template,context(),tmp_path/'never.docx',definition=definition)
    assert not (tmp_path/'never.docx').exists()


def test_unit_price_of_a_non_terminating_quotient_stays_within_the_context_schema():
    from decimal import Decimal
    from tidoc_print.context import _unit_price_text

    assert _unit_price_text(Decimal('100.00'), Decimal('3')) == '33.33333333'
    assert _unit_price_text(Decimal('356.50'), Decimal('3')) == '118.83333333'
    assert _unit_price_text(Decimal('988.00'), Decimal('2')) == '494.00'
    assert _unit_price_text(Decimal('1.00'), Decimal('8')) == '0.125'
    assert len(_unit_price_text(Decimal('1.00'), Decimal('7'))) <= 64

"""Real exports against synthetic, portable invoice fixtures."""
from copy import deepcopy
from decimal import Decimal
import json
import threading
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from tidoc.services.export_context import build_export_context, decimal_text, title_key, safe_name
from tidoc.services.exports import write_xlsx, write_attachment_zip


def raw(eid='a', amount='100.00', paid='90.00', title='测试大学', tax='TAX'):
    return {'id':eid,'title':title,'buyer_name':title,'buyer_tax_id':tax,'invoice_no':'000'+eid,'invoice_date':'2026-10-01','seller':'供应商<&>','profile_id':'person','fields':{'paid_amount':{'current':paid},'actual_item_name':{'current':'采购<&>名称'},'notes':{'current':'=1+1'}},'total':amount,'items':[{'name':'商品A','unit':'个','quantity':'2','total':'60'},{'name':'商品B','unit':'箱','quantity':'1','total':'40'}],'attachments':[]}


def test_decimal_missing_zero_negative_and_precision():
    context=build_export_context([raw('a','0','0'),raw('b','-0.123456789012345678901','-0.1')])
    assert context['totals']['invoice']=='-0.123456789012345678901'
    assert context['totals']['missing_paid']==0
    assert context['entries'][0]['paid_amount']=='0'
    assert decimal_text('NaN') is None
    context=build_export_context([raw('a','',None),raw('b','0','0')],output={'amount_basis':'paid'})
    assert context['totals']['invoice'] is None
    assert context['totals']['selected'] is None
    assert context['totals']['missing_paid']==1
    assert context['totals']['known_invoice']=='0'


def test_explicit_empty_default_outputs_stays_empty_in_resolver_and_planner(planner_setup):
    from tidoc.adapters.resolver import resolve_definition
    planner,entries,attachments,adapters,eid,root=planner_setup
    source={'scheme':{},'outputs':[{'id':'custom','default_selected':True}]}
    assert resolve_definition(source)['effective_settings']['print.default_outputs']==['custom']
    assert resolve_definition({**source,'scheme':{'default_outputs':[]}})['effective_settings']['print.default_outputs']==[]
    assert resolve_definition(source,{'print.default_outputs':[]})['effective_settings']['print.default_outputs']==[]
    original=adapters.get_scheme()
    revised=adapters.update_scheme_settings(original['id'],original['current_revision_id'],
                                           {'print.default_outputs':[]})
    with pytest.raises(ValueError,match='默认输出'):
        adapters.update_scheme_settings(revised['id'],revised['current_revision_id'],
                                       {'print.default_outputs':['not_declared']})
    person=entries.get(eid)['profile_id']
    new_id=entries.create(person)
    plan=planner.preview([new_id])
    assert plan['groups']==[]
    assert plan['ok'] is False
    assert any(item['code']=='NO_OUTPUT_SELECTED' for item in plan['diagnostics'])


def test_rows_preserve_units_and_invoice_paid_without_allocation():
    entry=raw()
    context=build_export_context([entry],output={'rows':'invoice_items','amount_basis':'paid'})
    assert [row['paid_amount'] for row in context['rows']]==['90.00',None]
    assert [row['total'] for row in context['rows']]==['60','40']
    assert context['totals']['paid']=='90.00'
    context=build_export_context([entry],output={'rows':'invoice_summary'})
    assert context['rows'][0]['quantity'] is None
    assert context['rows'][0]['unit'] is None
    context=build_export_context([{**entry,'items':[]}],output={'rows':'invoice_items'})
    assert context['rows'][0]['is_summary']
    assert context['rows'][0]['quantity'] is None
    assert context['rows'][0]['total']=='100.00'
    assert context['rows'][0]['product_name']=='采购<&>名称'


def test_xlsx_configured_columns_formula_text_and_unique_invoice_sum(tmp_path):
    context=build_export_context([raw()],output={'rows':'invoice_items'})
    output={'columns':[{'id':'notes','source':'entry.notes','label':'备注','width':20},{'id':'invoice','source':'entry.invoice.number','label':'票号'},{'id':'paid','source':'row.paid_amount','label':'实付','format':'money','total':'sum'},{'id':'exact','source':'entry.total','label':'精确值','format':'money','total':'sum'}]}
    path=write_xlsx(context,output,tmp_path/'out.xlsx')
    with zipfile.ZipFile(path) as z:
        xml=z.read('xl/worksheets/sheet1.xml')
        root=ET.fromstring(xml)
        assert b'<f>' not in xml
        assert b'=1+1' in xml
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert root.find('s:sheetViews/s:sheetView/s:pane',ns).get('state')=='frozen'
        assert root.find('s:autoFilter',ns) is not None
        assert root.find('.//s:c[@r="C4"]/s:v',ns).text=='90.00'
        assert root.find('.//s:c[@r="C3"]/s:v',ns) is None
        assert root.find('.//s:c[@r="B2"]/s:is/s:t',ns).text=='000a'


def test_zip_roles_subject_isolation_safe_names_and_private_payee(tmp_path):
    root=tmp_path/'resources';root.mkdir()
    (root/'one.txt').write_text('contract','utf-8')
    entries=[raw('a',title='同名',tax='1'),raw('b',title='同名',tax='2')]
    context=build_export_context(entries,payee={'name':'private','account_number':'SECRET'},output={'type':'attachment_zip'})
    resources=[{'id':e['id'],'entry_id':e['id'],'role_id':'custom:test:contract','path':'one.txt','original_name':'contract.txt'} for e in entries]
    output={'attachments':{'roles':['custom:test:contract'],'group_by':['title','entry'],'filename':'{entry.invoice_no}_{role.label}'}}
    path=write_attachment_zip(context,output,resources,tmp_path/'out.zip',root)
    with zipfile.ZipFile(path) as z:
        paths=[p for p in z.namelist() if p.endswith('.txt') and p!='清单.txt']
        assert len({p.split('/')[0] for p in paths})==2
        assert b'SECRET' not in z.read('manifest.json')
        assert 'custom:test:contract' in z.read('manifest.json').decode()
    assert safe_name('CON')=='_CON'
    with pytest.raises(ValueError):
        write_attachment_zip(context,output,[{**resources[0],'path':'../escape'}],tmp_path/'bad.zip',root)
    assert not (tmp_path/'bad.zip').exists()


@pytest.fixture
def planner_setup(tmp_path):
    from tidoc.db import Database,DataRoot,EntryRepo,ProfileRepo,AttachmentRepo
    from tidoc.adapters.service import AdapterService
    from tidoc.services.export_plan import ExportPlanner
    from tidoc.engine.models import ParsedInvoice,ParsedItem
    from reportlab.pdfgen import canvas
    root=DataRoot(tmp_path/'data');db=Database(root.db_path);adapters=AdapterService(db,root)
    adapters.bootstrap()
    generic=next(s for s in adapters.list_schemes() if s['package_id']=='org.tidoc.generic')
    adapters.complete_adapter_setup(generic['id'])
    profiles=ProfileRepo(db);entries=EntryRepo(db);attachments=AttachmentRepo(db,root)
    profile=profiles.create('测试甲','')
    parsed=ParsedInvoice(invoice_no='0001',buyer_name='测试大学',buyer_tax_id='TEST',total=Decimal('100.001'),items=[ParsedItem(name='零件',actual_name='零件',unit='个',quantity=Decimal('1'),total=Decimal('100.001'))])
    eid=entries.create(profile['id'],parsed=parsed)
    pdf=tmp_path/'invoice.pdf';c=canvas.Canvas(str(pdf));c.drawString(72,700,'INVOICE');c.save()
    attachments.add(eid,pdf,'invoice_pdf')
    entries.update_field(eid,'paid_amount','0')
    return ExportPlanner(db,root,adapters),entries,attachments,adapters,eid,root


def test_planner_real_all_four_outputs_and_saved_snapshot(planner_setup,tmp_path):
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['reimbursement','materials','generic_overview','generic_attachments'],{'date':'2026-10-06'})
    assert plan['ok'],plan['diagnostics']
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed',job['diagnostics']
    assert {f['type'] for f in job['files']}=={'docx','pdf_bundle','xlsx','attachment_zip'}
    assert all(Path(f['path']).exists() for f in job['files'])
    assert (root.export_jobs_dir/job['job_id']/'context.json').exists()
    before=deepcopy(job['snapshot'])
    entries.update_field(eid,'notes','now changed')
    recreated=planner.regenerate(job['job_id'])
    assert recreated['status']=='completed',recreated['diagnostics']
    assert recreated['snapshot']['files']==before['files']
    assert recreated['job_id']!=job['job_id']
    assert planner.get_job(job['job_id'])['snapshot']==before


def test_export_progress_and_cancel_remain_available_during_api_render(planner_setup,monkeypatch):
    from tidoc.api import Api
    from tidoc.services import export_plan
    planner,entries,attachments,adapters,eid,root=planner_setup
    api=Api(root.root)
    api._export_planner_instance=planner
    plan=planner.preview([eid],['generic_overview'],{'date':'2026-10-06'})
    assert plan['ok'],plan['diagnostics']
    started=threading.Event();proceed=threading.Event();results=[]
    original=export_plan.write_xlsx
    def slow_write(*args,**kwargs):
        started.set()
        if not proceed.wait(5):
            raise RuntimeError('test progress endpoint blocked')
        return original(*args,**kwargs)
    monkeypatch.setattr(export_plan,'write_xlsx',slow_write)
    worker=threading.Thread(target=lambda:results.append(api.run_export(plan['plan_id'])))
    worker.start()
    try:
        assert started.wait(3)
        progress=api.get_export_progress(plan['plan_id'])['data']
        assert progress['phase']=='rendering_core'
        assert progress['completed']==0 and progress['total']==1
        progress['completed']=999
        assert api.get_export_progress(plan['plan_id'])['data']['completed']==0
        assert api.cancel_export(plan['plan_id'])['ok']
    finally:
        proceed.set();worker.join(5)
    assert not worker.is_alive()
    assert results[0]['data']['status']=='cancelled'
    assert api.get_export_progress(plan['plan_id'])['data']['status']=='cancelled'
    assert not (root.exports_dir/plan['plan_id']).exists()
    api.db.close()


def test_cancellation_events_are_bounded_and_released_after_the_job(planner_setup):
    planner,entries,attachments,adapters,eid,root=planner_setup
    for index in range(300):
        planner.cancel(f'preview-{index}')   # cancelled previews never run
    assert len(planner._cancellations)<=256
    plan=planner.preview([eid],['generic_overview'],{'date':'2026-10-06'})
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed'
    assert job['id'] not in planner._cancellations


def test_initial_export_record_failure_is_recoverable_without_publishing(planner_setup,monkeypatch):
    from tidoc.services.export_plan import ExportPreflightError
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['generic_overview'],{'date':'2026-10-06'})
    original=planner._save_record_files
    def fail_record(job):
        directory=planner._job_dir(job['id']);directory.mkdir(parents=True,exist_ok=True)
        (directory/'context.json.tmp').write_text('partial record','utf-8')
        raise OSError('simulated disk full')
    monkeypatch.setattr(planner,'_save_record_files',fail_record)
    with pytest.raises(ExportPreflightError) as error:
        planner.run(plan['plan_id'])
    assert error.value.diagnostics[0]['code']=='EXPORT_RECORD_FAILED'
    job=planner.get_job(plan['plan_id'])
    assert job['status']=='failed' and job['snapshot']['entry_ids']==[eid]
    assert plan['plan_id'] not in planner._plans
    assert not planner._job_dir(job['id']).exists()
    assert not (root.exports_dir/job['id']).exists()
    assert entries.get(eid)['attachments']
    monkeypatch.setattr(planner,'_save_record_files',original)
    regenerated=planner.regenerate(job['id'])
    assert regenerated['status']=='completed'
    assert Path(regenerated['files'][0]['path']).is_file()
    assert planner.get_job(job['id'])['status']=='failed'


def test_planner_stale_missing_pdf_failure_atomic_and_cancel(planner_setup,tmp_path,monkeypatch):
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['generic_overview'])
    entries.update_field(eid,'notes','after preview')
    with pytest.raises(ValueError,match='重新预览'):
        planner.run(plan['plan_id'])
    plan=planner.preview([eid],['generic_overview','generic_attachments'])
    import tidoc.services.export_plan as module
    monkeypatch.setattr(module,'write_attachment_zip',lambda *a,**k:(_ for _ in ()).throw(OSError('disk full')))
    job=planner.run(plan['plan_id'])
    assert job['status']=='failed'
    assert not (root.exports_dir/job['id']).exists()
    monkeypatch.undo()
    plan=planner.preview([eid],['generic_overview'])
    planner.cancel(plan['plan_id'])
    with pytest.raises(ValueError):planner.run(plan['plan_id'])
    plan=planner.preview([eid],['materials'])
    assert plan['ok']
    resource=planner._plans[plan['plan_id']]['files'][0]['resources'][0]
    (root.root/resource['path']).unlink()
    job=planner.run(plan['plan_id'])
    assert job['status']=='failed'
    assert not (root.exports_dir/job['id']).exists()


def test_regenerate_rejects_changed_deleted_resource_and_preserves_original(planner_setup):
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['generic_attachments'])
    original=planner.run(plan['plan_id'])
    resource=original['resources'][0]
    (root.root/resource['path']).write_bytes(b'changed')
    regenerated=planner.regenerate(original['id'])
    assert regenerated['status']=='failed'
    assert Path(original['files'][0]['path']).exists()
    assert planner.get_job(original['id'])['status']=='completed'


def test_grouping_revisions_payees_same_name_tax_and_batches(planner_setup,tmp_path):
    from tidoc.adapters.resolver import resolve_definition
    from tidoc.db import ProfileRepo,BatchRepo
    from tidoc.engine.models import ParsedInvoice
    planner,entries,attachments,adapters,eid,root=planner_setup
    base=entries.get(eid);sid=base['scheme_id'];revision=base['scheme_revision_id']
    definition=deepcopy(adapters.get_revision(revision))
    definition['outputs']=[{'id':'custom_overview','label':'团队总览','type':'xlsx','default_selected':True,'payee_mode':'by_claimant','group_by':['title','batch'],'rows':'invoice_summary','amount_basis':'paid','filename':'{title.short_name}_{batch.name}.xlsx','columns':[{'id':'entry.invoice_no','label':'票号'},{'id':'entry.title.name','label':'抬头'},{'id':'entry.title.tax_id','label':'税号'},{'id':'row.paid_amount','label':'实付','format':'money','total':'sum'}]}]
    definition['fields']=[{'id':'signer','scope':'export','type':'text','label':'签署','required_at':['export'],'presentation':'visible','sensitive':False,'transfer':'never'}]
    from tidoc.adapters.registry import required_capabilities
    definition['manifest']['requires']['capabilities']=required_capabilities(definition)
    package=adapters.packages.revision_record(revision)['content_hash']
    new_revision=adapters.packages.store_revision(resolve_definition(definition),package)
    adapters.packages.link_revision(sid,new_revision)
    planner.db.conn.execute('UPDATE entries SET scheme_revision_id=? WHERE id=?',(new_revision,eid));planner.db.conn.commit()
    person=ProfileRepo(planner.db).create('测试乙','')
    eid2=entries.create(person['id'],parsed=ParsedInvoice(invoice_no='0002',buyer_name='测试大学',buyer_tax_id='DIFFERENT',total=Decimal('0')),scheme_id=sid,scheme_revision_id=new_revision)
    entries.update_field(eid2,'paid_amount','0')
    source=tmp_path/'second.xml';source.write_text('<invoice/>','utf-8')
    attachments.add(eid2,source,'invoice_xml')
    p1=adapters.payees.create(name='收款甲',personnel_id='001',bank_name='银行A',account_number='0001')
    p2=adapters.payees.create(name='收款乙',personnel_id='002',bank_name='银行B',account_number='0002')
    adapters.payees.set_mapping(sid,base['profile_id'],p1['id']);adapters.payees.set_mapping(sid,person['id'],p2['id'])
    BatchRepo(planner.db).create('批次A',entry_ids=[eid]);BatchRepo(planner.db).create('批次B',entry_ids=[eid2])
    plan=planner.preview([eid,eid2],['custom_overview'],{'date':'2026-10-06','export_fields_by_revision':{new_revision:{'signer':'测试签署'}}})
    assert plan['ok'],plan['diagnostics']
    assert len(plan['groups'])==2
    assert {g['payee']['account_tail'] for g in plan['groups']}=={'0001','0002'}
    assert {g['batch']['name'] for g in plan['groups']}=={'批次A','批次B'}
    snapshot=planner._plans[plan['plan_id']]
    assert all(f['context']['export']['fields']['signer']=='测试签署' for f in snapshot['files'])
    job=planner.run(plan['plan_id']);assert job['status']=='completed',job['diagnostics']
    assert len(job['files'])==2
    assert len({f['filename'] for f in job['files']})==2


def test_missing_payee_mapping_and_partial_default_no_field_fallback(planner_setup):
    planner,entries,attachments,adapters,eid,root=planner_setup
    base=entries.get(eid);sid=base['scheme_id']
    definition=deepcopy(adapters.get_revision(base['scheme_revision_id']))
    next(o for o in definition['outputs'] if o['id']=='reimbursement')['payee_mode']='single'
    revision=adapters.packages.store_revision(definition,adapters.packages.revision_record(base['scheme_revision_id'])['content_hash'])
    adapters.packages.link_revision(sid,revision)
    planner.db.conn.execute('UPDATE entries SET scheme_revision_id=? WHERE id=?',(revision,eid));planner.db.conn.commit()
    payee=adapters.payees.create(name='完整对象',bank_name='',account_number='0009')
    adapters.payees.set_default(sid,payee['id'])
    plan=planner.preview([eid],['reimbursement'])
    assert any(d['code']=='MISSING_PAYEE_FIELD' for d in plan['diagnostics'])
    plan=planner.preview([eid],['reimbursement'],{'payee_mode':'by_claimant'})
    assert any(d['code']=='MISSING_PAYEE' for d in plan['diagnostics'])


def test_fixed_output_settings_reject_runtime_and_batch_overrides(planner_setup):
    planner,*_=planner_setup
    definition={'scheme':{'settings':{'print.amount_basis':{'fixed':'invoice','editable':False},'print.numbering':{'fixed':True,'editable':False},'print.image_layout':{'default':'a4_landscape_2','editable':False}}},'effective_settings':{'print.amount_basis':'invoice','print.numbering':True,'print.image_layout':'a4_landscape_2'}}
    output={'id':'materials','type':'pdf_bundle','payee_mode':'none','amount_basis':'invoice','pdf':{}}
    for options in [{'amount_basis':'paid'},{'output_options':{'materials':{'pdf':{'numbering':False}}}},{'output_options':{'materials':{'pdf':{'image_layout':'a4_portrait_1'}}}},{'payee_mode':'single'},{'output_options':{'materials':{'attachments':{'roles':[]}}}}]:
        with pytest.raises(ValueError):planner._effective_output(output,definition,options)
    with pytest.raises(ValueError):planner._effective_output(output,definition,{}, {'print.numbering':False})
    assert planner._effective_output(output,definition,{'amount_basis':'invoice'})['pdf']['numbering'] is True
    assert planner._effective_output(output,{}, {'output_options':{'materials':{'pdf':{'content_order':'role'}}}})['pdf']['content_order']=='role'


def test_active_cancel_terminates_component_and_no_publication(planner_setup,monkeypatch):
    import tidoc.services.printing as printing
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['materials'])
    entered=threading.Event();finished=[]
    real_popen=printing.subprocess.Popen
    def slow_popen(cmd,**kwargs):
        import sys
        proc=real_popen([sys.executable,'-c','import time; time.sleep(30)'],**kwargs)
        entered.set()
        return proc
    monkeypatch.setattr(printing.subprocess,'Popen',slow_popen)
    thread=threading.Thread(target=lambda:finished.append(planner.run(plan['plan_id'])))
    thread.start();assert entered.wait(5)
    assert planner.cancel(plan['plan_id'])['status']=='cancel_requested'
    thread.join(5)
    assert not thread.is_alive()
    assert finished[0]['status']=='cancelled'
    assert not (root.exports_dir/plan['plan_id']).exists()
    assert not (root.export_jobs_dir/plan['plan_id']/'staging').exists()


def test_component_timeout_and_restart_recovery_do_not_publish(planner_setup,monkeypatch):
    import tidoc.services.printing as printing
    import sys
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['materials'],{'timeout_seconds':0.05})
    real_popen=printing.subprocess.Popen
    monkeypatch.setattr(printing.subprocess,'Popen',lambda cmd,**kwargs:real_popen([sys.executable,'-c','import time; time.sleep(30)'],**kwargs))
    result=planner.run(plan['plan_id'])
    assert result['status']=='failed'
    assert any('超时' in d['message'] for d in result['diagnostics'])
    assert not (root.exports_dir/plan['plan_id']).exists()
    snapshot=deepcopy(result['snapshot'])
    job=planner._new_job(snapshot)
    planner.jobs.update(job['id'],status='running')
    staging=root.export_jobs_dir/job['id']/'staging';staging.mkdir()
    (staging/'partial.xlsx').write_text('partial','utf-8')
    assert job['id'] in planner.recover_jobs()
    assert planner.jobs.get(job['id'])['status']=='failed'
    assert not staging.exists()


def test_planner_package_can_define_three_arbitrary_word_outputs(planner_setup,tmp_path):
    import shutil
    from docx import Document
    from tidoc.adapters.registry import required_capabilities
    planner,entries,attachments,adapters,eid,root=planner_setup
    source=tmp_path/'package'
    shutil.copytree(Path(__file__).resolve().parents[1]/'tidoc/builtin_adapters/org.tidoc.generic',source)
    shutil.rmtree(source/'templates');(source/'templates').mkdir()
    outputs=[]
    for oid,label in [('purchase_request','采购申请'),('settlement','结算说明'),('equipment_record','设备登记')]:
        document=Document();document.add_paragraph(label+' {{ title.name }} {{ export.output_id }} {{ totals.invoice|money }}')
        document.sections[0].header.paragraphs[0].text=label
        document.save(source/'templates'/f'{oid}.docx')
        outputs.append({'id':oid,'label':label,'type':'docx','default_selected':True,
                        'payee_mode':'none','group_by':['title'],'rows':'invoice_summary',
                        'amount_basis':'invoice','filename':label+'_{export.date}.docx',
                        'template':f'templates/{oid}.docx'})
    manifest=json.loads((source/'manifest.json').read_text('utf-8'));manifest['package_id']='org.tests.flexible-word'
    manifest['name']='研究团队单据';manifest['requires']['capabilities']=required_capabilities({'outputs':outputs})
    (source/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False),'utf-8')
    scheme=json.loads((source/'scheme.json').read_text('utf-8'));scheme['default_outputs']=[o['id'] for o in outputs]
    (source/'scheme.json').write_text(json.dumps(scheme,ensure_ascii=False),'utf-8')
    (source/'outputs.json').write_text(json.dumps({'outputs':outputs},ensure_ascii=False),'utf-8')
    if (source/'checksums.json').exists():(source/'checksums.json').unlink()
    preview=adapters.inspect_adapter(source)
    installed=adapters.install_adapter(preview['preview_id'])
    revision=installed['current_revision_id']
    planner.db.conn.execute('UPDATE entries SET scheme_id=?,scheme_revision_id=? WHERE id=?',(installed['id'],revision,eid))
    planner.db.conn.commit()
    plan=planner.preview([eid],options={'date':'2026-10-06'})
    assert plan['ok'],plan['diagnostics']
    assert {f['output_id'] for f in plan['files']}=={o['id'] for o in outputs}
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed',job['diagnostics']
    assert len(job['files'])==3
    for file in job['files']:
        declaration=next(o for o in outputs if o['id']==file['output_id'])
        document=Document(file['path'])
        assert document.paragraphs[0].text==declaration['label']+' 测试大学 '+declaration['id']+' 100.00'
        assert document.sections[0].header.paragraphs[0].text==declaration['label']
        assert Path(file['path']).name==declaration['label']+'_2026-10-06.docx'


def test_export_preflight_blocks_material_maximum_on_pinned_revision(planner_setup,tmp_path):
    from tidoc.adapters.registry import required_capabilities
    planner,entries,attachments,adapters,eid,root=planner_setup
    # Existing materials may exceed a newly selected revision even though the
    # attachment API correctly enforces the current limit for future additions.
    invoice_xml=tmp_path/'invoice.xml';invoice_xml.write_text('<invoice/>','utf-8')
    attachments.add(eid,invoice_xml,'invoice_xml')
    base=entries.get(eid);definition=deepcopy(adapters.get_revision(base['scheme_revision_id']))
    next(role for role in definition['materials'] if role['id']=='invoice')['max_count']=1
    definition['manifest']['requires']['capabilities']=required_capabilities(definition)
    revision=adapters.packages.store_revision(definition,adapters.packages.revision_record(base['scheme_revision_id'])['content_hash'])
    adapters.packages.link_revision(base['scheme_id'],revision)
    planner.db.conn.execute('UPDATE entries SET scheme_revision_id=? WHERE id=?',(revision,eid));planner.db.conn.commit()
    plan=planner.preview([eid],['overview'])
    assert not plan['ok']
    diagnostics=[d for d in plan['diagnostics'] if d['code']=='MATERIAL_MAXIMUM']
    assert any(d['stage']=='export' and d['entry_id']==eid and d['target']=='invoice' and d['output_id']=='overview' for d in diagnostics)
    with pytest.raises(ValueError):planner.run(plan['plan_id'])
    assert not list(root.exports_dir.iterdir())


def test_local_copies_sharing_revision_have_separate_output_and_field_choices(planner_setup,tmp_path):
    from tidoc.adapters.registry import required_capabilities
    from tidoc.engine.models import ParsedInvoice
    from jsonschema import Draft202012Validator
    planner,entries,attachments,adapters,eid,root=planner_setup
    original=entries.get(eid);sid=original['scheme_id']
    definition=deepcopy(adapters.get_revision(original['scheme_revision_id']))
    definition['fields']=[{'id':'signer','scope':'export','type':'text','label':'签署人',
                           'required_at':['export'],'presentation':'visible','sensitive':True,'transfer':'never'}]
    overview=next(o for o in definition['outputs'] if o['id']=='overview')
    overview['required_fields']=['export.fields.signer']
    overview['columns']=[{'id':'export.fields.signer','label':'签署人'}]
    definition['manifest']['requires']['capabilities']=required_capabilities(definition)
    revision=adapters.packages.store_revision(definition,adapters.packages.revision_record(original['scheme_revision_id'])['content_hash'])
    adapters.packages.set_current_revision(sid,revision)
    first=adapters.copy_scheme(sid,'研究组甲');second=adapters.copy_scheme(sid,'研究组乙')
    assert first['current_revision_id']==second['current_revision_id']==revision
    planner.db.conn.execute('UPDATE entries SET scheme_id=?,scheme_revision_id=? WHERE id=?',(first['id'],revision,eid));planner.db.conn.commit()
    eid2=entries.create(original['profile_id'],parsed=ParsedInvoice(invoice_no='COPY-002',buyer_name='测试大学',buyer_tax_id='TEST',total=Decimal('20')),scheme_id=second['id'],scheme_revision_id=revision)
    xml=tmp_path/'copy-invoice.xml';xml.write_text('<invoice/>','utf-8');attachments.add(eid2,xml,'invoice_xml')
    bindings=[first['id']+':'+revision,second['id']+':'+revision]
    fields={bindings[0]:{'signer':'甲私人签署'},bindings[1]:{'signer':'乙私人签署'}}
    options={'date':'2026-10-06','output_ids_by_binding':{bindings[0]:['overview'],bindings[1]:[]},
             'export_fields_by_binding':fields,'export_fields_by_revision':{revision:{'signer':'不应采用'}}}
    plan=planner.preview([eid,eid2],options=options)
    assert plan['ok'],plan['diagnostics']
    assert len(plan['files'])==1
    assert plan['groups'][0]['scheme_id']==first['id']
    assert plan['groups'][0]['entry_ids']==[eid]
    validator=Draft202012Validator(json.loads((Path(__file__).resolve().parents[1]/'schemas/team-adapter/1/context.schema.json').read_text('utf-8')))
    snapshot=planner._plans[plan['plan_id']]
    assert snapshot['files'][0]['context']['export']['fields']==fields[bindings[0]]
    assert not list(validator.iter_errors(snapshot['files'][0]['context']))
    assert '乙私人签署' not in json.dumps(snapshot['files'][0]['context'],ensure_ascii=False)
    options['output_ids_by_binding'][bindings[1]]=['overview']
    options['output_options_by_binding']={bindings[0]:{'overview':{'amount_basis':'paid'}},
                                          bindings[1]:{'overview':{'amount_basis':'invoice'}}}
    plan=planner.preview([eid,eid2],options=options)
    assert plan['ok'],plan['diagnostics']
    contexts={item['context']['scheme']['id']:item['context'] for item in planner._plans[plan['plan_id']]['files']}
    assert len(contexts)==2
    for scheme,binding,other in [(first,bindings[0],'乙私人签署'),(second,bindings[1],'甲私人签署')]:
        ctx=contexts[scheme['id']]
        assert ctx['export']['fields']==fields[binding]
        assert ctx['totals']['amount']==('0' if scheme['id']==first['id'] else '20')
        assert not list(validator.iter_errors(ctx))
        assert other not in json.dumps(ctx,ensure_ascii=False)
        assert not {'output_ids_by_binding','export_fields_by_binding','export_fields_by_revision','fields','payee_ids_by_binding','output_options_by_binding','_binding_key'} & ctx['export']['options'].keys()
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed',job['diagnostics']
    assert len(job['files'])==2
    for file in job['files']:
        with zipfile.ZipFile(file['path']) as archive:
            text=archive.read('xl/worksheets/sheet1.xml').decode()
        assert ('甲私人签署' in text) != ('乙私人签署' in text)


def test_new_plans_recover_old_missing_items_with_pinned_context_only(planner_setup,monkeypatch):
    from tidoc.engine import ParsedInvoice, ParsedItem
    planner,entries,attachments,adapters,eid,root=planner_setup
    planner.db.conn.execute('DELETE FROM items WHERE entry_id=?',(eid,))
    planner.db.conn.commit()
    entries.update_field(eid,'actual_item_name','人工核对名称')
    observed=[]
    def parsed(path,context=None):
        observed.append(context)
        return ParsedInvoice(invoice_no='0001',items=[ParsedItem(name='票面原名',actual_name='票面原名',unit='卷',quantity=Decimal('2'),total=Decimal('100.001'))])
    monkeypatch.setattr('tidoc.engine.parse_pdf',parsed)
    plan=planner.preview([eid],['reimbursement'])
    assert plan['ok'],plan['diagnostics']
    row=planner._plans[plan['plan_id']]['files'][0]['context']['rows'][0]
    assert row['actual_name']=='人工核对名称' and row['product_name']=='票面原名'
    assert row['unit']=='卷' and row['quantity']=='2'
    assert observed[0].titles==()
    assert entries.get(eid)['items']==[]
    monkeypatch.setattr('tidoc.engine.parse_pdf',lambda *a,**k:(_ for _ in ()).throw(ValueError('bad PDF')))
    plan=planner.preview([eid],['reimbursement'])
    assert plan['ok']
    assert planner._plans[plan['plan_id']]['files'][0]['context']['rows'][0]['actual_name']=='人工核对名称'


def test_new_imports_in_old_batch_use_frozen_defaults_and_xml_only(tmp_path):
    from tidoc.api import Api
    api=Api(tmp_path/'data')
    scheme=next(s for s in api.adapters.list_schemes() if s['package_id']=='org.tidoc.generic')
    api.complete_adapter_setup(scheme['id'])
    api.set_title_profiles([{'name':'旧单位','tax_id':'OLD'},{'name':'新单位','tax_id':'NEW'}])
    current=api.adapters.get_scheme()
    old_title=next(t['id'] for t in current['definition']['scheme']['titles'] if t['name']=='旧单位')
    new_title=next(t['id'] for t in current['definition']['scheme']['titles'] if t['name']=='新单位')
    old=api.adapters.update_scheme_settings(current['id'],current['current_revision_id'],{'entry.default_title_id':old_title,'entry.default_paid_to_invoice':False})
    batch=api.batches.create('保留旧规则')
    api.adapters.update_scheme_settings(old['id'],old['current_revision_id'],{'entry.default_title_id':new_title,'entry.default_paid_to_invoice':True})
    profile=api.profiles.create('测试人员','')
    paths=[]
    for number in ('OLD-001','OLD-002'):
        path=tmp_path/(number+'.xml')
        path.write_text(f'<root><EIid>{number}</EIid><BuyerName>旧单位</BuyerName><BuyerTaxID>OLD</BuyerTaxID><TotalTax-includedAmount>12.30</TotalTax-includedAmount></root>','utf-8')
        paths.append(path)
    first=api.create_entry(profile['id'],xml_path=str(paths[0]),batch_id=batch['id'])
    assert first['ok'],first
    second=api.batch_create_entries(profile['id'],[{'files':[{'path':str(paths[1]),'type':'invoice_xml'}]}],batch_id=batch['id'])
    assert second['ok'] and second['data']['created']==1,second
    for eid in [first['data']['id'],*second['data']['entry_ids']]:
        entry=api.entries.get(eid)
        assert entry['scheme_revision_id']==old['current_revision_id']
        assert entry['title']=='旧单位'
        assert entry['fields']['paid_amount']['current']==''
        assert api.batches.entry_ids(batch['id']).count(eid)==1
    response=api.preview_export(api.batches.entry_ids(batch['id']),['materials'])
    assert response['ok'],response  # 调用本身成功；阻断项在计划里，不能当成调用失败
    plan=response['data']
    assert not plan['ok'],plan
    assert any(d['code']=='MISSING_INVOICE_PDF' for d in plan['diagnostics'])


def test_export_destination_record_failure_keeps_recoverable_failed_job(planner_setup,monkeypatch):
    planner,entries,attachments,adapters,eid,root=planner_setup
    plan=planner.preview([eid],['overview'])
    write_text=Path.write_text
    def fail_destination(path,*args,**kwargs):
        if path.name=='destination.json':
            raise OSError('destination journal is not writable')
        return write_text(path,*args,**kwargs)
    monkeypatch.setattr(Path,'write_text',fail_destination)
    job=planner.run(plan['plan_id'])
    assert job['status']=='failed'
    assert planner.jobs.get(job['id'])['status']=='failed'
    assert planner.get_progress(job['id'])['status']=='failed'
    assert not list(root.exports_dir.iterdir())
    assert not (root.export_jobs_dir/job['id']/'staging').exists()
    assert entries.get(eid)['attachments']
    monkeypatch.setattr(Path,'write_text',write_text)
    assert planner.regenerate(job['id'])['status']=='completed'


def test_unbatched_export_fields_are_temporary_and_isolated(planner_setup):
    from tidoc.adapters.registry import required_capabilities
    planner,entries,attachments,adapters,eid,root=planner_setup
    entry=entries.get(eid);scheme_id=entry['scheme_id']
    definition=deepcopy(adapters.get_revision(entry['scheme_revision_id']))
    definition['fields']=[{'id':'project','scope':'batch','type':'text','label':'经费项目',
                           'required_at':['export'],'presentation':'visible','sensitive':False,'transfer':'never'}]
    output=next(o for o in definition['outputs'] if o['id']=='overview')
    output.update(required_fields=['batch.fields.project'],columns=[{'id':'batch.fields.project','label':'项目'}])
    definition['manifest']['requires']['capabilities']=required_capabilities(definition)
    revision=adapters.packages.store_revision(definition,adapters.packages.revision_record(entry['scheme_revision_id'])['content_hash'])
    adapters.packages.set_current_revision(scheme_id,revision)
    planner.db.conn.execute('UPDATE entries SET scheme_revision_id=? WHERE id=?',(revision,eid));planner.db.conn.commit()
    plan=planner.preview([eid],['overview'])
    assert not plan['ok']
    binding=scheme_id+':'+revision
    plan=planner.preview([eid],['overview'],{'batch_fields_by_binding':{binding:{'project':'本次经费'}}})
    assert plan['ok'],plan['diagnostics']
    snapshot=planner._plans[plan['plan_id']]
    ctx=snapshot['files'][0]['context']
    assert ctx['batch']['fields']['project']=='本次经费'
    assert ctx['entries'][0]['batch']['id']==''
    assert 'batch_fields_by_binding' not in ctx['export']['options']
    assert not entries.get(eid)['batches']
    assert planner.db.conn.execute('SELECT COUNT(*) FROM batches').fetchone()[0]==0
    assert planner.db.conn.execute('SELECT COUNT(*) FROM extension_values').fetchone()[0]==0
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed',job['diagnostics']
    assert planner.regenerate(job['id'])['status']=='completed'


def test_generic_material_pdf_supports_external_roles_and_title_isolation(planner_setup,tmp_path):
    from tidoc.engine import ParsedInvoice
    from pypdf import PdfReader
    planner,entries,attachments,adapters,eid,root=planner_setup
    original=entries.get(eid)
    eid2=entries.create(original['profile_id'],parsed=ParsedInvoice(invoice_no='GENERIC-002',buyer_name='另一主体',buyer_tax_id='SECOND',total=Decimal('9')))
    pdf=tmp_path/'external.pdf'
    from reportlab.pdfgen import canvas
    c=canvas.Canvas(str(pdf));c.drawString(72,700,'EXTERNAL');c.save()
    attachments.add(eid2,pdf,'invoice_pdf')
    # A received historical role can be printed by the core generic output.
    att=attachments.add(eid,pdf,'other')
    planner.db.conn.execute('UPDATE attachments SET role_id=? WHERE id=?',('custom:org.external:contract',att['id']))
    planner.db.conn.commit()
    plan=planner.preview([eid,eid2],['generic_materials'])
    assert plan['ok'],plan['diagnostics']
    assert len(plan['groups'])==2
    assert all(g['scheme_id'] is None for g in plan['groups'])
    assert {g['title']['name'] for g in plan['groups']}=={'测试大学','另一主体'}
    job=planner.run(plan['plan_id'])
    assert job['status']=='completed',job['diagnostics']
    assert sorted(len(PdfReader(file['path']).pages) for file in job['files'])==[1,2]
    non_pdf=tmp_path/'contract.txt';non_pdf.write_text('original contract','utf-8')
    attachments.add(eid,non_pdf,'other')
    blocked=planner.preview([eid],['generic_materials'])
    assert not blocked['ok']
    assert any(d['code']=='UNCONVERTIBLE_MATERIAL' for d in blocked['diagnostics'])


def test_xml_only_folder_import_and_pinned_suggested_tags(tmp_path):
    from tidoc.api import Api
    from tidoc.services.folder_import import scan_folder
    invoice=tmp_path/'only.xml'
    invoice.write_text('<root><EIid>98765432100123456789</EIid><BuyerName>样例单位</BuyerName><TotalTax-includedAmount>1.00</TotalTax-includedAmount></root>','utf-8')
    scanned=scan_folder(tmp_path)
    assert scanned['invoice_pdf_count']==0
    assert scanned['invoice_xml_only_count']==1
    assert len(scanned['groups'])==1 and not scanned['ungrouped']
    api=Api(tmp_path/'data');scheme=api.adapters.get_scheme()
    api.complete_adapter_setup(scheme['id'])
    api.adapters.update_scheme_settings(scheme['id'],scheme['current_revision_id'],{'entry.suggested_tags':['经费甲','耗材']})
    person=api.profiles.create('测试报账人','')
    result=api.batch_create_entries(person['id'],scanned['groups'])
    assert result['ok'] and result['data']['created']==1,result
    assert api.entries.get(result['data']['entry_ids'][0])['tags']==['经费甲','耗材']


def test_guard_hands_blocked_plans_to_the_frontend_as_data():
    # 预检计划带 ok 字段；有阻断项（ok=False、没有 error）时不能被当成调用失败。
    from tidoc.api import _guard

    class Bridge:
        _api_lock = __import__('threading').RLock()

        @_guard
        def blocked(self):
            return {'ok': False, 'plan_id': 'p', 'diagnostics': [{'code': 'X', 'severity': 'blocked', 'message': '缺少材料'}]}

        @_guard
        def passing(self):
            return {'ok': True, 'plan_id': 'p', 'diagnostics': []}

        @_guard
        def failing(self):
            return {'ok': False, 'error': '失败'}

    bridge = Bridge()
    blocked = bridge.blocked()
    assert blocked['ok'] is True and blocked['data']['ok'] is False and blocked['data']['diagnostics'][0]['message'] == '缺少材料'
    assert bridge.passing()['ok'] is True and bridge.passing()['plan_id'] == 'p'
    assert bridge.failing() == {'ok': False, 'error': '失败'}


def test_print_component_processes_never_open_a_console_window_on_windows(monkeypatch, tmp_path):
    # 打开设置时探测打印组件、以及生成时调用它，都是控制台程序；Windows 上必须隐藏黑色终端。
    import subprocess
    from types import SimpleNamespace
    from tidoc.services import printing

    seen = []

    def fake_run(cmd, **kwargs):
        seen.append(kwargs)
        return SimpleNamespace(returncode=0, stdout='{"component": "tidoc_print"}', stderr='')

    exe = tmp_path / 'tidoc_print.exe'
    exe.write_bytes(b'x')
    monkeypatch.setattr(printing.sys, 'platform', 'win32')
    monkeypatch.setattr(subprocess, 'CREATE_NO_WINDOW', 0x08000000, raising=False)
    monkeypatch.setattr(printing.subprocess, 'run', fake_run)
    printing._CAPABILITIES_CACHE.clear()
    printing._query_capabilities(exe)
    assert seen and seen[0]['creationflags'] == 0x08000000

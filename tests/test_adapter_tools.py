from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def run_cli(*args):
    return subprocess.run([sys.executable, "-m", "tidoc.adapter_tools", *map(str, args)], cwd=ROOT, text=True,encoding='utf-8', capture_output=True)


def test_cli_validate_json_and_diagnostic_exit(tmp_path):
    good = run_cli("validate", ROOT / "examples/adapters/generic", "--format", "json")
    assert good.returncode == 0, good.stderr
    assert json.loads(good.stdout)["ok"] is True
    invalid = tmp_path / "invalid"
    invalid.mkdir()
    (invalid / "manifest.json").write_text("{}", "utf-8")
    result = run_cli("validate", invalid, "--format", "json")
    assert result.returncode == 1
    assert json.loads(result.stdout or result.stderr)["ok"] is False


def test_init_diff_and_pack_are_real_package_operations(tmp_path):
    package = tmp_path / "team"
    init = run_cli("init", package, "--from", "generic", "--id", "org.example.team")
    assert init.returncode == 0, init.stderr
    manifest = json.loads((package / "manifest.json").read_text("utf-8"))
    assert manifest["package_id"] == "org.example.team"
    diff = run_cli("diff", ROOT / "examples/adapters/generic", package)
    assert diff.returncode == 0, diff.stderr
    assert json.loads(diff.stdout)["changed"] is True
    archive = tmp_path / "team.tidoc-preset"
    packed = run_cli("pack", package, "--out", archive)
    assert packed.returncode == 0, packed.stderr
    assert archive.is_file()
    with zipfile.ZipFile(archive) as bundle:
        assert "checksums.json" in bundle.namelist()
        assert bundle.testzip() is None
    verified = run_cli("validate", archive, "--format", "json")
    assert verified.returncode == 0, verified.stderr


def test_explain_and_render_use_fixture_and_actual_writer(tmp_path):
    package = ROOT / "examples/adapters/generic"
    fixture = ROOT / "examples/adapters/fixtures/generic.json"
    explained = run_cli("explain", package, "--fixture", fixture)
    assert explained.returncode == 0, explained.stderr
    data = json.loads(explained.stdout)
    assert data["fixture_entries"] == 2
    assert "settings" in data and "grouping" in data
    rendered = run_cli("render", package, "--fixture", fixture, "--out", tmp_path, "--output", "reimbursement")
    assert rendered.returncode == 0, rendered.stderr
    outputs = json.loads(rendered.stdout)["files"]
    assert outputs and all(Path(item["path"]).is_file() and item["bytes"] > 0 for item in outputs)


def test_case_runner_evaluates_declared_expectations(tmp_path):
    cases = tmp_path / "cases"
    cases.mkdir()
    spec = {"schema_version": 1, "entries": [{"invoice_no": "CASE-1", "total": "9.20", "materials": ["invoice", "payment_screenshot"]}],
            "expect": {"ok": True, "group_counts": {"reimbursement": 1}}}
    (cases / "basic.json").write_text(json.dumps(spec), "utf-8")
    result = run_cli("test", ROOT / "examples/adapters/generic", "--cases", cases)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["cases"][0]["ok"] is True


def test_every_example_fixture_checks_expectations_and_renders(tmp_path):
    for name in ("generic", "lab", "club", "minimal", "bitfsae"):
        result = run_cli("test", ROOT / "examples/adapters" / name, "--cases", ROOT / "examples/adapters/fixtures" / f"{name}.json")
        assert result.returncode == 0, f"{name}: {result.stdout}\n{result.stderr}"
        report = json.loads(result.stdout)
        assert report["ok"] is True
        assert report["cases"][0]["ok"] is True
        assert report["cases"][0]["files"]


def test_lab_fixture_omitting_conditionally_required_approval_fails(tmp_path):
    fixture = json.loads((ROOT / "examples/adapters/fixtures/lab.json").read_text("utf-8"))
    fixture["entries"][0]["materials"].remove("custom:org.example.lab:approval")
    missing = tmp_path / "lab-without-approval.json"
    missing.write_text(json.dumps(fixture, ensure_ascii=False), "utf-8")

    result = run_cli("render", ROOT / "examples/adapters/lab", "--fixture", missing,
                     "--out", tmp_path / "rendered", "--output", "cost_settlement")
    assert result.returncode == 1
    diagnostic = json.loads(result.stdout)["errors"][0]
    assert diagnostic["code"] == "RULE_MATERIAL_REQUIRED"
    assert diagnostic["target"] == "custom:org.example.lab:approval"


def test_lab_arbitrary_word_id_by_claimant_flows_through_render_test_and_pack(tmp_path):
    package = ROOT / "examples/adapters/lab"
    fixture = ROOT / "examples/adapters/fixtures/lab.json"

    rendered = run_cli("render", package, "--fixture", fixture, "--out", tmp_path / "rendered", "--output", "cost_settlement")
    assert rendered.returncode == 0, rendered.stdout + rendered.stderr
    render_result = json.loads(rendered.stdout)
    assert render_result["ok"] is True
    assert render_result["outputs"] == ["cost_settlement"]
    assert render_result["output_stats"]["cost_settlement"]["payee_mode"] == "by_claimant"
    claimant_groups = [group for group in render_result["groups"] if group["output_id"] == "cost_settlement"]
    assert len(claimant_groups) == 2
    assert all(Path(item["path"]).is_file() for item in render_result["files"])
    assert _docx_contains(render_result["files"], "实验室费用结算")

    tested = run_cli("test", package, "--cases", fixture)
    assert tested.returncode == 0, tested.stdout + tested.stderr
    case_result = json.loads(tested.stdout)["cases"][0]
    assert case_result["ok"] is True, case_result["failures"]
    assert case_result["files"]

    archive = tmp_path / "lab.tidoc-preset"
    packed = run_cli("pack", package, "--out", archive)
    assert packed.returncode == 0, packed.stdout + packed.stderr
    with zipfile.ZipFile(archive) as bundle:
        definition = json.loads(bundle.read("outputs.json"))
        output = next(item for item in definition["outputs"] if item["id"] == "cost_settlement")
        assert output["label"] == "费用结算"
        assert output["template"] == "templates/cost_settlement.docx"
        assert "templates/cost_settlement.docx" in bundle.namelist()


def test_init_rewrites_lab_and_club_material_namespaces(tmp_path):
    for source in ("lab", "club"):
        package_id = f"org.example.copied.{source}"
        target = tmp_path / source
        created = run_cli("init", target, "--from", source, "--id", package_id)
        assert created.returncode == 0, created.stderr
        validation = run_cli("validate", target, "--format", "json")
        assert validation.returncode == 0, validation.stdout or validation.stderr
        assert json.loads(validation.stdout)["ok"] is True
        for path in target.rglob("*.json"):
            text = path.read_text("utf-8")
            assert f"custom:org.example.{source}:" not in text
        materials = json.loads((target / "materials.json").read_text("utf-8"))["roles"]
        assert any(role["id"].startswith(f"custom:{package_id}:") for role in materials)
        rules = json.loads((target / "rules.json").read_text("utf-8"))["rules"]
        references = [requirement.get("material") for rule in rules for requirement in rule.get("require", [])]
        assert any(value and value.startswith(f"custom:{package_id}:") for value in references)


def test_pack_renders_all_docx_outputs_even_when_condition_is_false(tmp_path, monkeypatch):
    from tidoc.adapter_tools import cli

    source = ROOT / "examples/adapters/minimal"
    copied = tmp_path / "conditional"
    shutil.copytree(source, copied)
    outputs_path = copied / "outputs.json"
    outputs = json.loads(outputs_path.read_text("utf-8"))
    conditional = dict(outputs["outputs"][0])
    conditional["id"] = "conditional_docx"
    conditional["when"] = {"field": "invoice.total", "op": "eq", "value": "999.00"}
    outputs["outputs"].append(conditional)
    outputs_path.write_text(json.dumps(outputs, ensure_ascii=False, indent=2), "utf-8")
    manifest_path = copied / "manifest.json"
    manifest = json.loads(manifest_path.read_text("utf-8"))
    manifest["requires"]["capabilities"] = sorted(set(manifest["requires"]["capabilities"]) | {"rules.v1"})
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), "utf-8")

    real_render = cli._render
    rendered_ids = []
    staged_conditions = []

    def capture_render(package_path, fixture_path, out_dir, output_ids=None):
        rendered_ids.extend(output_ids or [])
        stage_outputs = json.loads((Path(package_path) / "outputs.json").read_text("utf-8"))
        staged_conditions.extend(item.get("when") for item in stage_outputs["outputs"] if item["id"] in output_ids)
        return real_render(package_path, fixture_path, out_dir, output_ids)

    monkeypatch.setattr(cli, "_render", capture_render)
    archive = tmp_path / "conditional.tidoc-preset"
    cli._strict_pack(copied, archive)
    assert rendered_ids == ["reimbursement", "conditional_docx"]
    assert staged_conditions == [None, None]
    with zipfile.ZipFile(archive) as bundle:
        published = json.loads(bundle.read("outputs.json"))
    assert published["outputs"][1]["when"] == conditional["when"]


def test_pack_runs_real_renderer_for_each_example_with_docx(tmp_path):
    for name in ("generic", "lab", "club", "minimal", "bitfsae"):
        archive = tmp_path / f"{name}.tidoc-preset"
        result = run_cli("pack", ROOT / "examples/adapters" / name, "--out", archive)
        assert result.returncode == 0, f"{name}: {result.stdout}\n{result.stderr}"
        with zipfile.ZipFile(archive) as bundle:
            assert bundle.testzip() is None
            assert "checksums.json" in bundle.namelist()


def test_reference_is_generated_from_runtime_registry(tmp_path):
    target = tmp_path / "REFERENCE.md"
    result = run_cli("reference", "--out", target)
    assert result.returncode == 0, result.stderr
    text = target.read_text("utf-8")
    assert "entry.default_paid_to_invoice" in text
    assert "payee.account_number" in text


def test_packaged_registry_catalogs_match_runtime_definitions():
    from tidoc.adapters.registry import FIELD_CATALOG, FILENAME_FIELDS, SETTINGS

    catalog = json.loads((ROOT / "schemas/team-adapter/1/registry.json").read_text("utf-8"))
    assert catalog["field_catalog"] == FIELD_CATALOG
    assert catalog["settings"] == SETTINGS
    assert catalog["filename_fields"] == list(FILENAME_FIELDS)


def _docx_contains(files, fragment):
    from docx import Document

    for item in files:
        if item["type"] != "docx":
            continue
        document = Document(item["path"])
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        if fragment in text:
            return True
    return False

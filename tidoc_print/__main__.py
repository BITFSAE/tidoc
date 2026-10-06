"""Bounded JSON IPC entrypoint for the optional component."""
from __future__ import annotations
import argparse
import json
import os
from pathlib import Path
from .protocol import capabilities, render_request


def self_test_status():
    status = capabilities()
    result = {'ok': not status['missing'], **status}
    if not result['ok']:
        return result
    try:
        from tempfile import TemporaryDirectory
        from docx import Document
        from PIL import Image
        from pypdf import PdfReader, PdfWriter
        from .context import build_export_context
        with TemporaryDirectory(prefix='tidoc-print-self-test-') as directory:
            root = Path(directory) / 'resources'; root.mkdir()
            template = root / 'template.docx'
            document = Document(); document.add_paragraph('{{ title.name }} {{ totals.invoice|money(true) }}')
            document.save(template)
            writer = PdfWriter(); writer.add_blank_page(300, 400); writer.write(root / 'invoice.pdf')
            Image.new('RGB', (40, 60), 'blue').save(root / 'payment.png')
            entry = {'id': 'self-test', 'title': '自检主体', 'total': '0'}
            files = []
            for kind, filename in (('docx', 'document.docx'), ('pdf_bundle', 'materials.pdf')):
                output = {'id': 'self_test_' + kind, 'type': kind, 'payee_mode': 'none',
                          'pdf': {'include_roles': ['invoice', 'payment_screenshot'],
                                  'image_layout': 'a4_landscape_2', 'numbering': True}}
                item = {'output': output, 'filename': filename,
                        'context': build_export_context([entry], output=output)}
                if kind == 'docx':
                    item['template'] = template.name
                else:
                    item['resources'] = [
                        {'id': 'invoice', 'entry_id': entry['id'], 'role_id': 'invoice', 'path': 'invoice.pdf'},
                        {'id': 'payment', 'entry_id': entry['id'], 'role_id': 'payment_screenshot', 'path': 'payment.png'}]
                files.append(item)
            rendered = render_request({'ipc_version': 2, 'job_id': 'component-self-test',
                'resources_root': str(root), 'output_dir': str(Path(directory) / 'output'), 'files': files})
            if Document(rendered['files'][0]['path']).paragraphs[0].text != '自检主体 ¥0.00':
                raise ValueError('Word 内容检查失败')
            if len(PdfReader(rendered['files'][1]['path']).pages) != 2:
                raise ValueError('PDF 页数检查失败')
            result['smoke_outputs'] = ['docx', 'pdf_bundle']
    except Exception as exc:
        result.update(ok=False, code='COMPONENT_SELF_TEST_FAILED', error=str(exc))
    return result


def main():
    parser = argparse.ArgumentParser(prog='tidoc_print')
    parser.add_argument('--input')
    parser.add_argument('--result')
    parser.add_argument('--capabilities', action='store_true')
    parser.add_argument('--self-test', action='store_true')
    parser.add_argument('--timeout',type=float,default=120)
    parser.add_argument('--cancel-file')
    args = parser.parse_args()
    if args.capabilities:
        print(json.dumps(capabilities(),ensure_ascii=False)); return 0
    if args.self_test:
        status = self_test_status()
        print(json.dumps(status,ensure_ascii=False))
        return 0 if status['ok'] else 1
    if not args.input or not args.result:
        parser.error('--input 和 --result 必须同时提供')
    result_path = Path(args.result)
    try:
        if Path(args.input).stat().st_size>32*1024*1024:
            raise ValueError('打印请求文件超过限制')
        payload=json.loads(Path(args.input).read_text('utf-8'))
        payload['timeout_seconds']=min(args.timeout,payload.get('timeout_seconds',120))
        if args.cancel_file:
            payload['cancel_file']=args.cancel_file
        data=render_request(payload)
        _write_result(result_path,{'ok':True,'data':data}); return 0
    except Exception as exc:
        _write_result(result_path,{'ok':False,'error':str(exc),'code':getattr(exc,'code','RENDER_FAILED'),'diagnostics':getattr(exc,'diagnostics',[])})
        return 1


def _write_result(path,result):
    path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(result,ensure_ascii=False,indent=2),'utf-8')
    os.replace(tmp,path)


if __name__=='__main__':
    raise SystemExit(main())

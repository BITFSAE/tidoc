"""Validate, explain, render, test, compare and package team adapters."""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
import hashlib
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from ..adapters.loader import AdapterValidationError, load_package, pack_package, validate_package
from ..adapters.policy import evaluate_condition
from ..adapters.registry import CAPABILITIES, FIELD_CATALOG, SETTINGS


class EnvironmentFailure(RuntimeError):
    pass


def _json(path):
    try:
        return json.loads(Path(path).read_text("utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取 JSON 文件 {path}：{exc}") from exc


def _fixture(path):
    value = _json(path)
    if not isinstance(value, dict) or value.get("schema_version") != 1 or not isinstance(value.get("entries"), list):
        raise ValueError("fixture 必须包含 schema_version: 1 和 entries 数组。")
    if not value["entries"]:
        raise ValueError("fixture 至少要有一条发票记录。")
    return value


def _replace_strings(value, replacements):
    if isinstance(value, str):
        for old, new in replacements:
            value = value.replace(old, new)
        return value
    if isinstance(value, list):
        return [_replace_strings(item, replacements) for item in value]
    if isinstance(value, dict):
        return {key: _replace_strings(item, replacements) for key, item in value.items()}
    return value
def _default_value(field):
    if "default" in field:
        return field["default"]
    kinds = {"text": "Fixture 值", "multiline": "Fixture 值", "integer": 0, "decimal": "0", "money": "0.00", "date": "2026-10-06", "boolean": False}
    if field.get("type") in ("select", "multiselect"):
        options = field.get("options", [])
        if field["type"] == "multiselect":
            return [options[0]["value"]] if options else []
        return options[0]["value"] if options else ""
    return kinds.get(field.get("type"), "")


def _field_values(definition, scope, supplied):
    result = {field["id"]: _default_value(field) for field in definition.get("fields", []) if field["scope"] == scope}
    if supplied:
        result.update(supplied)
    return result


def _make_file(kind, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if kind == "pdf":
        from reportlab.pdfgen import canvas
        doc = canvas.Canvas(str(path), pagesize=(595, 842))
        doc.setFont("Helvetica", 12)
        doc.drawString(54, 790, f"Tidoc synthetic fixture document {path.name}")
        doc.drawString(54, 768, "Fictional data for adapter validation")
        doc.save()
    elif kind == "image":
        from PIL import Image, ImageDraw
        image = Image.new("RGB", (800, 600), "white")
        ImageDraw.Draw(image).text((40, 40), f"Tidoc fixture image {path.name}", fill="black")
        image.save(path, "PNG")
    elif kind == "text":
        path.write_text(f"Tidoc synthetic fixture material {path.name}\n", "utf-8")
    else:
        raise ValueError(f"fixture 不支持生成此文件类型：{kind}")
    return path


def _fixture_materials(source, definition, fixture_root, generated_root):
    supplied = list(source.get("attachments") or [])
    for role_id in source.get("materials", []):
        if not any(item.get("role_id") == role_id for item in supplied):
            supplied.append({"role_id": role_id, "generate": "auto"})
    normalized = []
    for index, spec in enumerate(supplied):
        if not isinstance(spec, dict) or not spec.get("role_id"):
            raise ValueError("fixture 附件必须声明 role_id。")
        role_id = spec["role_id"]
        extension = spec.get("extension")
        generate = spec.get("generate")
        if generate == "auto" or not generate and not spec.get("path"):
            role = next((item for item in definition.get("materials", []) if item["id"] == role_id), {})
            extensions = role.get("extensions") or [".pdf"]
            extension = extension or extensions[0].lower()
            generate = "image" if extension in (".png", ".jpg", ".jpeg") else "text" if extension == ".txt" else "pdf"
        if generate:
            kind = generate
            extension = extension or {"pdf": ".pdf", "image": ".png", "text": ".txt"}.get(kind)
            if not extension:
                raise ValueError(f"fixture 生成类型无效：{kind}")
            source_path = _make_file(kind, generated_root / f"material-{index:03}{extension}")
        else:
            relative = Path(spec.get("path", ""))
            if not str(relative) or relative.is_absolute() or ".." in relative.parts:
                raise ValueError("fixture 附件 path 必须是 fixture 目录内的相对路径。")
            source_path = (fixture_root / relative).resolve()
            if not source_path.is_relative_to(fixture_root.resolve()) or not source_path.is_file():
                raise ValueError(f"fixture 附件不存在或越界：{relative}")
        normalized.append({**spec, "source_path": source_path, "role_id": role_id})
    return normalized


def _create_fixture_entries(api, definition, fixture_data, fixture_root):
    from ..db.attachments import TYPE_INSPECTION, TYPE_INVOICE_PDF, TYPE_OTHER, TYPE_PAYMENT, TYPE_PHYSICAL_IMAGE
    from ..engine.models import ParsedInvoice, ParsedItem
    import tempfile as _tempfile

    entry_ids = []
    batch_specs = {}
    with _tempfile.TemporaryDirectory(prefix="tidoc-fixture-materials-") as generated:
        for index, source in enumerate(fixture_data["entries"], 1):
            total = source.get("total")
            try:
                amount = Decimal(str(total)) if total not in (None, "") else Decimal("0")
            except (InvalidOperation, ValueError):
                amount = Decimal("0")
            items = [ParsedItem(name=str(item.get("name", item.get("actual_name", "Fixture 商品"))),
                                actual_name=str(item.get("actual_name", item.get("name", "Fixture 商品"))),
                                unit=str(item.get("unit", "个")),
                                quantity=Decimal(str(item.get("quantity", "1"))) if item.get("quantity", "1") is not None else None,
                                total=Decimal(str(item.get("total", total or "0"))), spec=str(item.get("spec", "")))
                     for item in source.get("items", [])]
            title = source.get("buyer_name") or source.get("title") or (definition.get("scheme", {}).get("titles") or [{}])[0].get("name", "示例单位")
            parsed = ParsedInvoice(invoice_no=str(source.get("invoice_no") or f"FIXTURE-{index:03}"),
                                   invoice_date=str(source.get("invoice_date") or fixture_data.get("export_date", "2026-10-06")),
                                   seller=str(source.get("seller") or "虚构示例供应商"), buyer_name=title,
                                   buyer_tax_id=str(source.get("buyer_tax_id", "")), total=amount, items=items,
                                   source="fixture", total_present=total not in (None, ""))
            claimant_spec = source.get("claimant") or {}
            claimant = claimant_spec.get("name", f"示例报账人{index}") if isinstance(claimant_spec, dict) else str(claimant_spec)
            reviewer = source.get("reviewer") or (claimant_spec.get("reviewer", "") if isinstance(claimant_spec, dict) else "")
            if not reviewer and definition.get("effective_settings", {}).get("profile.reviewer_required"):
                reviewer = "示例审核人"
            profile = api.profiles.create(claimant, reviewer)
            eid = api.entries.create(profile["id"], parsed=parsed)
            entry_ids.append(eid)
            paid = source.get("paid_amount", total)
            if "paid_amount" in source or paid is not None:
                api.entries.update_field(eid, "paid_amount", "" if paid is None else str(paid))
            for key, value in _field_values(definition, "entry", source.get("fields") or {}).items():
                api.adapters.save_extension_values("entry", eid, {key: value}, api.entries.get(eid).get("scheme_id"))
            fixture_payees = {item.get("id"): item for item in fixture_data.get("payees", []) if isinstance(item, dict)}
            payee_values = source.get("payee") or fixture_payees.get(source.get("payee_id")) or fixture_data.get("payee")
            if payee_values and "id" in payee_values:
                payee_values = {key: value for key, value in payee_values.items() if key != "id"}
            if any(output.get("payee_mode") in ("single", "by_claimant") for output in definition.get("outputs", [])):
                payee_values = payee_values or {"name": claimant, "personnel_id": f"F{index:04}", "contact": "000-0000", "account_type": "personal_bank", "bank_name": "示例银行", "account_number": f"000000000000{index:04}"}
            if payee_values:
                generated_payee = {"name": claimant, "personnel_id": f"F{index:04}", "contact": "000-0000",
                                   "account_type": "personal_bank", "bank_name": "示例银行",
                                   "account_number": f"000000000000{index:04}"}
                generated_payee.update(payee_values)
                payee = api.adapters.save_payee(values=generated_payee)
                api.adapters.set_payee_mapping(api.adapters.get_scheme()["id"], profile["id"], payee["id"])
                if not api.adapters.payees.get_default(api.adapters.get_scheme()["id"]):
                    api.adapters.set_scheme_payee(api.adapters.get_scheme()["id"], payee["id"])
                for key, value in _field_values(definition, "payee", source.get("payee_fields") or {}).items():
                    api.adapters.save_extension_values("payee", payee["id"], {key: value}, api.adapters.get_scheme()["id"])
            materials = _fixture_materials(source, definition, fixture_root, Path(generated))
            for material in materials:
                role = material["role_id"]
                kind = {"invoice": TYPE_INVOICE_PDF, "payment_screenshot": TYPE_PAYMENT, "physical_image": TYPE_PHYSICAL_IMAGE, "inspection_pdf": TYPE_INSPECTION}.get(role, TYPE_OTHER)
                api.attachments.add(eid, material["source_path"], kind, role_id=role)
            batch = source.get("batch")
            if batch:
                batch_name = batch if isinstance(batch, str) else batch.get("name", "Fixture 批次")
                batch_specs.setdefault(batch_name, {"entries": [], "values": (batch.get("fields", {}) if isinstance(batch, dict) else {})})["entries"].append(eid)
        for name, spec in batch_specs.items():
            batch = api.batches.create(name, entry_ids=spec["entries"])
            values = _field_values(definition, "batch", spec["values"])
            if values:
                api.adapters.save_extension_values("batch", batch["id"], values, api.adapters.get_scheme()["id"])
    for key, value in _field_values(definition, "scheme", fixture_data.get("scheme_fields") or {}).items():
        api.adapters.save_extension_values("scheme", api.adapters.get_scheme()["id"], {key: value}, api.adapters.get_scheme()["id"])
    return entry_ids


def _fixture_runtime(package_path, fixture_path):
    package = load_package(package_path)
    fixture_data = _fixture(fixture_path)
    temp = tempfile.TemporaryDirectory(prefix="tidoc-adapter-cli-")
    try:
        from ..api import Api
        api = Api(Path(temp.name) / "data")
        preview = api.adapters.inspect_adapter(package_path)
        scheme = api.adapters.install_adapter(preview["preview_id"], {"mode": "install", "name": package.name, "set_default": True})
        api.adapters.complete_adapter_setup(scheme["id"], {})
        definition = api.adapters.get_scheme(scheme["id"])["definition"]
        entry_ids = _create_fixture_entries(api, definition, fixture_data, Path(fixture_path).resolve().parent)
        return temp, api, package, definition, fixture_data, entry_ids
    except BaseException:
        temp.cleanup()
        raise


def _all_output_ids(definition):
    return [output["id"] for output in definition.get("outputs", [])] + ["generic_overview", "generic_attachments"]


def _preview(package_path, fixture_path, output_ids=None):
    temp, api, package, definition, fixture_data, entry_ids = _fixture_runtime(package_path, fixture_path)
    try:
        output_ids = output_ids or _all_output_ids(definition)
        options = {"date": fixture_data.get("export_date", "2026-10-06"), "fields": _field_values(definition, "export", fixture_data.get("export_fields") or {})}
        plan = api._export_planner().preview(entry_ids, output_ids, options)
        saved = api._export_planner()._plans.get(plan["plan_id"], {})
        return temp, api, package, definition, fixture_data, entry_ids, plan, saved
    except BaseException:
        api.db.close()
        temp.cleanup()
        raise


def _explain(package_path, fixture_path):
    temp, api, package, definition, fixture_data, entry_ids, plan, saved = _preview(package_path, fixture_path)
    try:
        rule_results = []
        contexts = [item["context"] for item in saved.get("files", [])]
        for rule in definition.get("rules", []):
            values = []
            for context in contexts:
                for entry in context.get("entries", []):
                    rule_context = {**context, "entry": entry, "invoice": entry.get("invoice", {}), "definition": definition}
                    value = evaluate_condition(rule.get("when", {}), rule_context, stage=rule.get("stage", "complete"), definition=definition)
                    values.append("unknown" if value is None else value)
            rule_results.append({"rule_id": rule.get("id"), "results": values, "message": rule.get("message", ""), "requires": rule.get("require", [])})
        settings_sources = {}
        for key, value in definition.get("effective_settings", {}).items():
            descriptor = definition.get("scheme", {}).get("settings", {}).get(key, {})
            settings_sources[key] = {"value": value, "source": "fixed" if "fixed" in descriptor else "package" if "default" in descriptor else "core"}
        return {"ok": plan["ok"], "package_id": package.package_id, "package_version": package.package_version,
                "settings": settings_sources, "rules": rule_results, "diagnostics": plan["diagnostics"],
                "groups": plan["groups"],
                "grouping": {output_id: sum(group.get("output_id") == output_id for group in plan["groups"])
                             for output_id in dict.fromkeys(group.get("output_id") for group in plan["groups"])},
                "files": plan["files"], "fixture_entries": len(entry_ids)}
    finally:
        api.db.close()
        temp.cleanup()


def _render(package_path, fixture_path, out_dir, output_ids=None):
    temp, api, package, definition, fixture_data, entry_ids, plan, saved = _preview(package_path, fixture_path, output_ids)
    try:
        if not plan["ok"]:
            return {"ok": False, "fixture": str(fixture_path), "errors": plan["diagnostics"], "files": []}
        result = api._export_planner().run(plan["plan_id"], out_dir)
        if result.get("status") != "completed":
            return {"ok": False, "fixture": str(fixture_path), "errors": result.get("diagnostics", []), "files": []}
        contexts = saved.get("files", [])
        output_stats = {}
        all_entries = {}
        for item in contexts:
            context = item.get("context", {})
            output = item.get("output", {})
            output_id = output.get("id")
            stat = output_stats.setdefault(output_id, {"row_count": 0, "entries": {}, "payee_mode": output.get("payee_mode", "none")})
            stat["row_count"] += len(context.get("rows", []))
            for entry in context.get("entries", []):
                stat["entries"][entry.get("id")] = entry
                all_entries[entry.get("id")] = entry
        for stat in output_stats.values():
            entries = list(stat["entries"].values())
            invoice = [entry.get("total") for entry in entries]
            paid = [entry.get("paid_amount") for entry in entries]
            stat["invoice_total"] = None if any(value is None for value in invoice) else format(sum((Decimal(str(value)) for value in invoice), Decimal(0)), "f")
            stat["paid_total"] = None if any(value is None for value in paid) else format(sum((Decimal(str(value)) for value in paid), Decimal(0)), "f")
            stat["entry_count"] = len(entries)
        material_roles = {eid: sorted({att.get("role_id") for att in api.attachments.list(eid)}) for eid in entry_ids}
        distinct_titles = {(entry.get("title", {}).get("name", ""), entry.get("title", {}).get("tax_id", "")) for entry in all_entries.values()}
        groups = result.get("snapshot", {}).get("groups", [])
        rendered_files = [{**item, "bytes": Path(item["path"]).stat().st_size} for item in result.get("files", [])]
        return {"ok": True, "fixture": str(fixture_path), "job_id": result["job_id"],
                "groups": groups, "files": rendered_files, "output_dir": result.get("output_dir"),
                "outputs": sorted(output_stats), "output_stats": output_stats,
                "entry_count": len(entry_ids), "title_groups": len(distinct_titles), "material_roles": material_roles}
    finally:
        api.db.close()
        temp.cleanup()


def _test(package_path, cases_dir):
    root = Path(cases_dir)
    if not root.exists():
        raise ValueError(f"样例目录不存在：{root}")
    cases = [root] if root.is_file() else sorted(root.glob("*.json"))
    if not cases:
        raise ValueError(f"样例目录没有 JSON 样例：{root}")
    results = []
    for case in cases:
        spec = _json(case)
        fixture = (case.parent / spec.get("fixture", case.name)).resolve()
        if not fixture.is_file():
            fixture = case
        expectations = spec.get("expected", spec.get("expect", {}))
        failures = []
        output_ids = expectations.get("outputs") or spec.get("outputs") or _all_output_ids(load_package(package_path).definition)
        render_temp = tempfile.mkdtemp(prefix="tidoc-adapter-case-")
        rendered = _render(package_path, fixture, Path(render_temp) / "rendered", output_ids)
        actual_ok = rendered["ok"]
        expected_ok = expectations.get("ok", True)
        if actual_ok != expected_ok:
            failures.append(f"ok 期望 {expected_ok}，实际 {actual_ok}")
        diagnostics = rendered.get("errors", [])
        wanted = set(expectations.get("diagnostic_codes", []))
        actual_codes = {item["code"] for item in diagnostics}
        if not wanted <= actual_codes:
            failures.append("缺少诊断：" + ", ".join(sorted(wanted - actual_codes)))
        if actual_ok:
            if "entry_count" in expectations and rendered["entry_count"] != expectations["entry_count"]:
                failures.append(f"条目数期望 {expectations['entry_count']}，实际 {rendered['entry_count']}")
            if "invoice_total" in expectations and not _same_amount(expectations["invoice_total"], _stat_amount(rendered, output_ids, "invoice_total")):
                failures.append(f"发票金额期望 {expectations['invoice_total']}，实际 {_stat_amount(rendered, output_ids, 'invoice_total')}")
            if "paid_total" in expectations and not _same_amount(expectations["paid_total"], _stat_amount(rendered, output_ids, "paid_total")):
                failures.append(f"实付金额期望 {expectations['paid_total']}，实际 {_stat_amount(rendered, output_ids, 'paid_total')}")
            if "title_groups" in expectations and rendered["title_groups"] != expectations["title_groups"]:
                failures.append(f"抬头组数期望 {expectations['title_groups']}，实际 {rendered['title_groups']}")
            if "summary_row_count" in expectations:
                stats = rendered["output_stats"]
                row_count = next((stats[item]["row_count"] for item in output_ids if item in stats and item in ("overview", "generic_overview")), None)
                if row_count is None:
                    row_count = next((stats[item]["row_count"] for item in output_ids if item in stats), 0)
                if row_count != expectations["summary_row_count"]:
                    failures.append(f"汇总行数期望 {expectations['summary_row_count']}，实际 {row_count}")
            if "by_claimant_groups" in expectations:
                claimant_outputs = {key for key, stat in rendered["output_stats"].items()
                                    if stat["payee_mode"] == "by_claimant"}
                count = sum(group.get("output_id") in claimant_outputs for group in rendered.get("groups", []))
                if count != expectations["by_claimant_groups"]:
                    failures.append(f"按报账人分组数期望 {expectations['by_claimant_groups']}，实际 {count}")
            if "outputs" in expectations:
                missing = set(expectations["outputs"]) - set(rendered["outputs"])
                if missing:
                    failures.append("缺少实际输出：" + ", ".join(sorted(missing)))
            if "required_materials" in expectations:
                required = set(expectations["required_materials"])
                for entry_id, roles in rendered["material_roles"].items():
                    missing = required - set(roles)
                    if missing:
                        failures.append(f"条目 {entry_id} 缺少材料角色：" + ", ".join(sorted(missing)))
                        break
            file_count = len(rendered["files"])
            if "file_count" in expectations and file_count != expectations["file_count"]:
                failures.append(f"文件数期望 {expectations['file_count']}，实际 {file_count}")
            wanted_types = set(expectations.get("output_types", []))
            got_types = {item["type"] for item in rendered["files"]}
            if not wanted_types <= got_types:
                failures.append("缺少输出类型：" + ", ".join(sorted(wanted_types - got_types)))
            for output_id, count in expectations.get("group_counts", {}).items():
                groups = [group for group in rendered.get("groups", []) if group.get("output_id") == output_id]
                if len(groups) != count:
                    failures.append(f"{output_id} 分组数期望 {count}，实际 {len(groups)}")
            for fragment in expectations.get("contains_text", []):
                if not _outputs_contain(rendered["files"], fragment):
                    failures.append(f"输出未包含文本：{fragment}")
            for fragment in expectations.get("excludes_text", []):
                if _outputs_contain(rendered["files"], fragment):
                    failures.append(f"输出不应包含文本：{fragment}")
        results.append({"case": case.name, "ok": not failures, "failures": failures, "files": rendered.get("files", []), "diagnostics": diagnostics})
        shutil.rmtree(render_temp, ignore_errors=True)
    return {"ok": all(item["ok"] for item in results), "cases": results}


def _same_amount(expected, actual):
    if expected is None or actual is None:
        return expected is actual
    try:
        return Decimal(str(expected)) == Decimal(str(actual))
    except InvalidOperation:
        return False


def _stat_amount(rendered, output_ids, key):
    stats = rendered.get("output_stats", {})
    preferred = next((stats[output_id] for output_id in output_ids if output_id in stats and output_id in ("overview", "generic_overview")), None)
    preferred = preferred or next((stats[output_id] for output_id in output_ids if output_id in stats), {})
    return preferred.get(key)


def _outputs_contain(files, fragment):
    from zipfile import ZipFile
    for item in files:
        path = Path(item["path"])
        try:
            if item["type"] == "docx":
                from docx import Document
                document = Document(path)
                text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
            elif item["type"] == "pdf_bundle":
                from pypdf import PdfReader
                text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
            else:
                with ZipFile(path) as archive:
                    text = "\n".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.endswith((".xml", ".json", ".txt")))
            if fragment in text:
                return True
        except (OSError, ValueError, KeyError):
            continue
    return False


def _diff(old_path, new_path):
    old, new = load_package(old_path), load_package(new_path)
    a, b = old.definition, new.definition
    sections = {}
    for name in ("manifest", "scheme", "fields", "materials", "rules", "outputs"):
        if a.get(name) != b.get(name):
            sections[name] = {"before": a.get(name), "after": b.get(name)}
    return {"changed": bool(sections), "package": {"before": old.package_id, "after": new.package_id},
            "versions": {"before": old.package_version, "after": new.package_version},
            "content_hash": {"before": old.content_hash, "after": new.content_hash}, "changes": sections}


def _strict_pack(source, out):
    package = load_package(source)
    docx_outputs = [o for o in package.definition.get("outputs", []) if o.get("type") == "docx"]
    if docx_outputs:
        try:
            import docxtpl  # noqa: F401
            import jinja2  # noqa: F401
        except ImportError as exc:
            raise EnvironmentFailure("pack 含 DOCX 模板，必须安装 requirements-print.txt 并完成真实渲染验证。") from exc
        with tempfile.TemporaryDirectory(prefix="tidoc-adapter-pack-") as temp:
            stage = Path(temp) / "package"
            stage.mkdir()
            for name, data in package.files.items():
                if name == "checksums.json":
                    continue
                target = stage / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
            output_file = stage / "outputs.json"
            output_data = _json(output_file)
            # `when` controls whether a user may select an output. During pack,
            # every DOCX must still reach the common renderer at least once.
            for output in output_data.get("outputs", []):
                if output.get("type") == "docx":
                    output.pop("when", None)
                    output.pop("applies_when", None)
            output_file.write_text(json.dumps(output_data, ensure_ascii=False, indent=2) + "\n", "utf-8")
            fixture = Path(temp) / "fixture.json"
            # Pack validation uses an explicit synthetic fixture that exercises
            # every declared role without mutating user fixture input from tests.
            fixture_materials = [role["id"] for role in package.definition.get("materials", [])]
            fixture.write_text(json.dumps({"schema_version": 1, "export_date": "2026-10-06", "entries": [{"invoice_no": "FIXTURE-001", "total": "12.50", "materials": fixture_materials}]}, ensure_ascii=False), "utf-8")
            checked = _render(stage, fixture, Path(temp) / "render", [output["id"] for output in docx_outputs])
            if not checked["ok"]:
                raise AdapterValidationError(checked.get("errors") or [{"code": "TEMPLATE_RENDER_FAILED", "file": "", "location": "/", "message": "DOCX 模板未能通过真实渲染。"}])
    return pack_package(source, out)


def parser():
    root = argparse.ArgumentParser(prog="python -m tidoc.adapter_tools", description="校验、解释、渲染和打包 Tidoc 团队适配包。")
    sub = root.add_subparsers(dest="command", required=True)
    init = sub.add_parser("init", help="从内置包复制完整源目录")
    init.add_argument("directory"); init.add_argument("--from", dest="source", choices=("generic", "bitfsae", "minimal", "lab", "club"), default="generic"); init.add_argument("--id", required=True)
    for command in ("validate",):
        p = sub.add_parser(command); p.add_argument("package"); p.add_argument("--format", choices=("text", "json"), default="text")
    p = sub.add_parser("explain"); p.add_argument("package"); p.add_argument("--fixture", required=True); p.add_argument("--format", choices=("text", "json"), default="text")
    p = sub.add_parser("render"); p.add_argument("package"); p.add_argument("--fixture", required=True); p.add_argument("--out", required=True); p.add_argument("--output", action="append")
    p = sub.add_parser("test"); p.add_argument("package"); p.add_argument("--cases", required=True)
    p = sub.add_parser("diff"); p.add_argument("old"); p.add_argument("new")
    p = sub.add_parser("pack"); p.add_argument("package"); p.add_argument("--out", required=True)
    p = sub.add_parser("reference", help="从当前注册表生成开发参考页"); p.add_argument("--out", default="docs/adapters/REFERENCE.md")
    return root


def main(argv=None):
    args = parser().parse_args(argv)
    try:
        if args.command == "init":
            target = Path(args.directory)
            if target.exists() and any(target.iterdir()):
                raise ValueError(f"目标目录非空：{target}")
            source = Path(__file__).resolve().parents[2] / "examples" / "adapters" / args.source
            if not source.is_dir():
                raise EnvironmentFailure(f"内置示例缺失：{source}")
            target.mkdir(parents=True, exist_ok=True)
            shutil.copytree(source, target, dirs_exist_ok=True)
            manifest_path = target / "manifest.json"
            manifest = _json(manifest_path)
            old_id = manifest["package_id"]
            if not re.fullmatch(r"[a-z][a-z0-9_-]*(?:\.[a-z][a-z0-9_-]*)+", args.id) or len(args.id) > 100:
                raise ValueError("--id 必须是长度不超过 100 的小写点分包 ID。")
            manifest["package_id"] = args.id
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", "utf-8")
            replacements = [(f"custom:{old_id}:", f"custom:{args.id}:")]
            for json_path in target.rglob("*.json"):
                data = _json(json_path)
                data = _replace_strings(data, replacements)
                json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", "utf-8")
            load_package(target)
            print(json.dumps({"ok": True, "directory": str(target), "from": args.source, "package_id": args.id}, ensure_ascii=False))
            return 0
        if args.command == "validate":
            result = validate_package(args.package)
        elif args.command == "explain":
            result = _explain(args.package, args.fixture)
        elif args.command == "render":
            result = _render(args.package, args.fixture, args.out, args.output)
        elif args.command == "test":
            result = _test(args.package, args.cases)
        elif args.command == "diff":
            result = _diff(args.old, args.new)
        elif args.command == "pack":
            result = {"ok": True, "package": str(_strict_pack(args.package, args.out))}
        else:
            target = Path(args.out)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(_reference_markdown(), "utf-8")
            result = {"ok": True, "path": str(target)}
        if getattr(args, "format", "text") == "json":
            print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        else:
            print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        return 0 if result.get("ok", True) else 1
    except EnvironmentFailure as exc:
        print(json.dumps({"ok": False, "errors": [{"code": "ENVIRONMENT_UNAVAILABLE", "file": "", "location": "/", "message": str(exc)}], "warnings": []}, ensure_ascii=False), file=sys.stderr)
        return 2
    except AdapterValidationError as exc:
        stream = sys.stdout if getattr(args, "format", "text") == "json" else sys.stderr
        print(json.dumps({"ok": False, "errors": exc.diagnostics, "warnings": []}, ensure_ascii=False, indent=2), file=stream)
        return 1
    except (ValueError, OSError, KeyError, TypeError) as exc:
        stream = sys.stdout if getattr(args, "format", "text") == "json" else sys.stderr
        print(json.dumps({"ok": False, "errors": [{"code": "ADAPTER_TOOL_FAILED", "file": "", "location": "/", "message": str(exc)}], "warnings": []}, ensure_ascii=False, indent=2), file=stream)
        return 1


def _reference_markdown():
    lines = ["# 注册表参考", "", "本页由 `python -m tidoc.adapter_tools reference` 生成。设置由 `tidoc/adapters/registry.py` 定义；上下文字段由核心与组件共用 `tidoc_print/context_fields.json` 定义。", "", "## 设置", "", "| 设置 ID | 类型 | 核心默认值 | 允许作用域 |", "| --- | --- | --- | --- |"]
    for key, spec in SETTINGS.items():
        default = json.dumps(spec.get("default"), ensure_ascii=False)
        lines.append(f"| `{key}` | `{spec['type']}` | `{default}` | {', '.join(spec.get('scopes', []))} |")
    lines += ["", "## 模板与规则上下文字段", "", "| 字段 | 类型 |", "| --- | --- |"]
    for key, kind in FIELD_CATALOG.items():
        lines.append(f"| `{key}` | `{kind}` |")
    lines += ["", "## 能力 ID", "", *[f"- `{item}`" for item in CAPABILITIES], ""]
    return "\n".join(lines)
